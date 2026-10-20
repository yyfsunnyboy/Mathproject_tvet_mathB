"""Scoped replacement guard for unpublished general textbook imports."""

from __future__ import annotations

import hashlib
import io
import json
import sqlite3
from pathlib import Path


def inspect_general_pair(docx_upload, pdf_upload, volume: str) -> dict:
    """Read the uploaded pair without importing it or trusting its filename."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from core.textbook_processor import normalize_general_high_chapter_section, scan_general_high_layout

    docx_bytes, pdf_bytes = docx_upload.read(), pdf_upload.read()
    docx_upload.seek(0)
    pdf_upload.seek(0)
    if not docx_bytes.startswith(b"PK") or not pdf_bytes.startswith(b"%PDF"):
        raise ValueError("invalid_general_source_pair")
    doc = Document(io.BytesIO(docx_bytes))
    chunks = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            value = Paragraph(child, doc).text.strip()
            if value:
                chunks.append(value)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, doc).rows:
                chunks.extend(cell.text.strip() for cell in row.cells if cell.text.strip())
    scan = scan_general_high_layout("\n".join(chunks))
    unit = scan.get("chapter") or {}
    chapter, _ = normalize_general_high_chapter_section(f"{unit.get('unit', '')} {unit.get('title', '')}")
    if not unit.get("unit") or not unit.get("title") or not scan.get("items") or not volume:
        raise ValueError("general_source_chapter_unresolved")
    return {
        "curriculum": "general", "volume": volume, "chapter": chapter,
        "question_count": len(scan["items"]),
        "docx_sha256": hashlib.sha256(docx_bytes).hexdigest(),
        "pdf_sha256": hashlib.sha256(pdf_bytes).hexdigest(),
    }


def backup_sqlite_database(database_path: Path, backup_path: Path) -> None:
    """SQLite online backup, including committed WAL data."""
    database_path, backup_path = Path(database_path), Path(backup_path)
    backup_path.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(f"file:{database_path.resolve().as_posix()}?mode=ro", uri=True) as source:
        with sqlite3.connect(backup_path) as target:
            source.backup(target)
            if target.execute("pragma quick_check").fetchone()[0] != "ok":
                raise ValueError("replacement_backup_invalid")


def scope_fingerprint(rows) -> str:
    data = [(row.id, row.skill_id, row.source_description, row.problem_text, row.notes) for row in rows]
    return hashlib.sha256(json.dumps(data, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def validate_replacement_scope(rows, *, expected_ids: list[int], incoming_sha256: str) -> None:
    """Fail closed when the chapter has changed or contains independent sources."""
    import json

    if [row.id for row in rows] != expected_ids:
        raise ValueError("general_replacement_scope_changed")
    sources = set()
    for row in rows:
        provenance = json.loads(row.notes or "{}").get("source_provenance") or {}
        primary = (provenance.get("primary") or {}).get("sha256")
        pdf = (provenance.get("optional_enrich_pdf") or {}).get("sha256")
        sources.add((primary, pdf))
    if len(sources) > 1:
        raise ValueError("general_replacement_multiple_sources")


def replace_old_rows_in_session(session, model, policy: dict, curriculum_info: dict, parsed_data: dict) -> int:
    """Delete only the confirmed old rows inside the import's still-open transaction."""
    from sqlalchemy import text
    from core.textbook_processor import normalize_general_high_chapter_section

    preview = policy["preview"]
    if curriculum_info.get("curriculum") != "general" or str(curriculum_info.get("volume")) != preview["volume"]:
        raise ValueError("general_replacement_scope_mismatch")
    chapters = [normalize_general_high_chapter_section(str(ch.get("chapter_title") or ""))[0]
                for ch in parsed_data.get("chapters", []) if isinstance(ch, dict)]
    if chapters != [preview["chapter"]]:
        raise ValueError("general_replacement_chapter_mismatch")
    rows = session.query(model).filter_by(
        source_curriculum="general", source_volume=preview["volume"],
        source_chapter=preview["chapter"],
    ).order_by(model.id).all()
    validate_replacement_scope(rows, expected_ids=policy["old_ids"], incoming_sha256=preview["docx_sha256"])
    if policy.get("old_rows_fingerprint") and scope_fingerprint(rows) != policy["old_rows_fingerprint"]:
        raise ValueError("general_replacement_scope_changed")
    ids = [row.id for row in rows]
    if not ids:
        raise ValueError("general_replacement_empty_old_scope")
    skills = sorted({row.skill_id for row in rows})
    checks = (
        ("gencode_component_tracker", "textbook_example_id", ids),
        ("quiz_attempts", "question_id", ids),
        ("practice_attempts", "question_uid", [str(i) for i in ids]),
        ("mistake_notebook_entries", "source_question_uid", [str(i) for i in ids]),
        ("practice_attempts", "skill_id", skills),
        ("mistake_notebook_entries", "skill_id", skills),
        ("mistake_logs", "skill_id", skills),
        ("progress", "skill_id", skills),
    )
    for table, column, values in checks:
        if not values:
            continue
        params = {f"v{i}": value for i, value in enumerate(values)}
        placeholders = ",".join(f":v{i}" for i in range(len(values)))
        count = session.execute(text(f'SELECT COUNT(*) FROM "{table}" WHERE "{column}" IN ({placeholders})'), params).scalar_one()
        if count:
            raise ValueError(f"general_replacement_relation_blocked:{table}.{column}:{count}")
    deleted = session.query(model).filter(model.id.in_(ids)).delete(synchronize_session=False)
    if deleted != len(ids):
        raise ValueError("general_replacement_delete_count_mismatch")
    session.flush()
    return deleted


_STALE_BINDINGS_SQL = (
    'SELECT sc.id, sc.skill_id FROM skill_curriculum sc '
    "WHERE sc.curriculum = 'general' AND sc.volume = :volume AND sc.chapter = :chapter "
    'AND NOT EXISTS (SELECT 1 FROM textbook_examples te WHERE te.skill_id = sc.skill_id '
    "AND te.source_curriculum = 'general' AND te.source_volume = :volume AND te.source_chapter = :chapter) "
    'ORDER BY sc.id'
)


def stale_chapter_skill_bindings(session, volume: str, chapter: str) -> list[tuple[int, str]]:
    """``skill_curriculum`` rows of one 普高 unit whose skill has no textbook example in that unit.

    Scope is curriculum + volume + chapter only; skills themselves are never touched.
    """
    from sqlalchemy import text

    rows = session.execute(text(_STALE_BINDINGS_SQL), {"volume": str(volume), "chapter": str(chapter)}).all()
    return [(int(row[0]), str(row[1])) for row in rows]


def cleanup_stale_chapter_skill_bindings(session, policy: dict) -> int:
    """Delete this unit's invalid bindings inside the import transaction (idempotent).

    Call after the new unit rows are flushed so skills used by the new material are kept.
    """
    from sqlalchemy import text

    preview = policy["preview"]
    stale = stale_chapter_skill_bindings(session, preview["volume"], preview["chapter"])
    if not stale:
        return 0
    deleted = 0
    for binding_id, _skill_id in stale:
        deleted += session.execute(
            text("DELETE FROM skill_curriculum WHERE id = :id AND curriculum = 'general' "
                 "AND volume = :volume AND chapter = :chapter"),
            {"id": binding_id, "volume": preview["volume"], "chapter": preview["chapter"]},
        ).rowcount
    if deleted != len(stale):
        raise ValueError("general_replacement_binding_delete_count_mismatch")
    session.flush()
    return deleted
