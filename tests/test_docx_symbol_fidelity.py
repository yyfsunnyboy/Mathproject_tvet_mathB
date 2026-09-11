from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest

from core.textbook_processor_v2 import _docx_paragraph_text_with_symbols, phase1_extract_docx_lines


@pytest.mark.parametrize('font,code,expected', [
    ('Symbol', 'F02D', '−'), ('Symbol', '002D', '−'),
    ('Symbol', 'F02B', '+'), ('Wingdings', 'F02D', '[MATH_PARSE_FAILED:symbol:Wingdings:F02D]'),
])
def test_font_aware_symbol_order_and_source_unchanged(font, code, expected, tmp_path):
    doc = Document()
    p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    p.add_run('angle(')
    sym = OxmlElement('w:sym')
    sym.set(qn('w:font'), font)
    sym.set(qn('w:char'), code)
    p.add_run()._r.append(sym)
    p.add_run('73°)')
    original = p._p.xml
    assert _docx_paragraph_text_with_symbols(p) == f'angle({expected}73°)'
    assert p._p.xml == original
    path = tmp_path/'symbols.docx'
    doc.save(path)
    assert phase1_extract_docx_lines(str(path)) == [f'angle({expected}73°)']
