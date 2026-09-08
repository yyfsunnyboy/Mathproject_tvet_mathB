"""B2 1-1 source regressions; all DB operations use a seeded in-memory database."""

from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import copy
import json
import re

from docx import Document
from flask import Flask
import pytest

from models import db, SkillInfo, SkillCurriculum, TextbookExample
import core.textbook_processor_v2 as tp
import core.textbook_importer_v3_pipeline as pipeline
from core.textbook_b2_11 import SDG_TITLE, SOURCE_STEM, is_b2_11, correct_pdf_regions
from core.textbook_mathtype_converter import convert_docx_mathtype_to_latex_docx
from core.textbook_question_anchor import build_anchors_from_block_meta
from core.textbook_pdf_visual import enrich_textbook_examples_with_pdf_visuals

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "textbook_import/source/vocational/math_B2" / (SOURCE_STEM + ".docx")
PDF = SOURCE.with_suffix(".pdf")
SKILLS = {
    "DirectedAngle": "有向角",
    "AngleMeasurementAndConversion": "角的度量與換算",
    "ArcLengthAndAreaOfSector": "扇形的弧長與面積",
    "CoterminalAngles": "同界角",
}


@pytest.fixture
def section(tmp_path, monkeypatch):
    app = Flask(__name__)
    app.config.update(SQLALCHEMY_DATABASE_URI="sqlite:///:memory:",
                      SQLALCHEMY_TRACK_MODIFICATIONS=False)
    db.init_app(app)
    monkeypatch.setattr(tp, "_DOCX_BLOCK_META", {})
    monkeypatch.setattr(tp, "_MATHB_SECTION_CONCEPTS", {})
    monkeypatch.setattr(tp, "get_model", lambda *a, **k: pytest.fail("Unexpected AI call"))
    info = dict(curriculum="vocational", volume="數學B2", grade=10,
                chapter="第1章 三角函數", section="1-1 角度的基本性質",
                section_code="1-1", chapter_index=1, source_scope="section_textbook",
                original_filename=SOURCE.name, parse_filename=SOURCE.name)
    with app.app_context():
        db.create_all()
        for en, zh in {"outline_vocational_數學B2_11": None, **SKILLS}.items():
            sid = en if zh is None else "vh_數學B2_" + en
            db.session.add(SkillInfo(skill_id=sid, skill_en_name=en, skill_ch_name=zh or "outline",
                                    description="fixture", gemini_prompt="fixture"))
            db.session.add(SkillCurriculum(skill_id=sid, curriculum=info["curriculum"],
                                           volume=info["volume"], grade=10,
                                           chapter=info["chapter"], section=info["section"],
                                           paragraph=zh))
        db.session.commit()
        yield app, info
        db.session.remove()


def extract(tmp_path, info):
    if not SOURCE.is_file():
        pytest.skip("B2 1-1 source missing")
    converted = tmp_path / (SOURCE_STEM + "_Latex.docx")
    report = convert_docx_mathtype_to_latex_docx(SOURCE, converted)
    assert report["converted_ok"] == 142 and report["converted_failed"] == 0
    lines = tp.phase1_extract_docx_lines(str(converted), curriculum_info=info)
    blocks = tp.phase2_deterministic_block_slice(lines, curriculum_info=info)
    return converted, blocks


def test_chapter_uses_exact_existing_outline(section):
    _, info = section
    incoming = {**info, "chapter": "1", "section": ""}
    resolved = pipeline._fill_chapter_section_from_outline_or_lines(incoming, [])
    assert (resolved["chapter"], resolved["section"]) == (info["chapter"], info["section"])


def test_source_questions_table_and_phase3_existing_skills(section, tmp_path):
    _, info = section
    converted, blocks = extract(tmp_path, info)
    keys = list(blocks)
    assert len(keys) == 20
    assert keys[3:6] == ["隨堂練習2", SDG_TITLE, "例3"]
    sdg = tp._DOCX_BLOCK_META[SDG_TITLE]
    assert re.findall(r"\(([12])\)", sdg["problem_text"]) == ["1", "2"]
    assert re.findall(r"\(([12])\)", sdg["detailed_solution"]) == ["1", "2"]
    assert "SDGS" not in sdg["detailed_solution"]
    table = [t for t in Document(converted).tables
             if len(t.columns) == 12 and t.cell(0, 0).text == "度"][-1]
    text = blocks["1-1習題 基礎題 1"]
    rows = [ln for ln in text.splitlines() if ln.startswith("|")]
    assert len(rows) == 3
    for emitted, source in zip((rows[0], rows[2]), table.rows):
        assert [c.strip() for c in emitted.split("|")[1:-1]] == [
            tp._normalize_docx_line_text(c.text) for c in source.cells]
    for title, tokens in {
        "例1": ["135°", r"-\frac { 5 } { 3 }\pi"],
        "隨堂練習1": [r"-570{}^\circ", r"\frac { 5 } { 6 }\pi"],
        "1-1習題 基礎題 2": ["450", "560"],
        "1-1習題 基礎題 3": [r"1=\frac { 180{}^\circ } { \pi }", r"\frac { 6 } { 5 }\pi", r"\frac { 5 } { 3 }\pi", "4="],
    }.items():
        positions = [blocks[title].index(t) for t in tokens]
        assert positions == sorted(positions)
    parsed = tp.phase3_ai_metadata_alignment(sorted(blocks), info, None)
    items = [item for con in parsed["chapters"][0]["sections"][0]["concepts"]
             for bucket in ("examples", "practice_questions") for item in con.get(bucket, [])]
    assert [item["title"] for item in items] == keys
    for title, meta in tp._DOCX_BLOCK_META.items():
        if "基礎題" in title:
            n = int(title.split()[-1])
            expected = "AngleMeasurementAndConversion" if n <= 3 else (
                "ArcLengthAndAreaOfSector" if n <= 5 else "CoterminalAngles")
            assert meta["formal_skill_id"] == "vh_數學B2_" + expected
        assert meta["formal_skill_id"] in {"vh_數學B2_" + k for k in SKILLS}
        assert "MATH_PARSE_FAILED" not in meta["problem_text"] + meta["detailed_solution"]
        binding = tp._phase4_resolve_mathb_formal_binding(
            block_meta=meta, source_type=meta["source_type"], db_problem_text=meta["problem_text"],
            curriculum_info=info, item_sec_code="1-1", coords=tp._import_scope_coords(info))
        assert binding[1] == meta["formal_skill_id"]
        assert binding[2].chapter == info["chapter"]
    assert SkillInfo.query.count() == 5 and SkillCurriculum.query.count() == 5
    assert TextbookExample.query.count() == 0


def test_missing_existing_skill_stops_without_creation(section, tmp_path):
    _, info = section
    _, blocks = extract(tmp_path, info)
    SkillCurriculum.query.filter_by(paragraph="角的度量與換算").delete()
    db.session.flush()
    with pytest.raises(ValueError, match="existing skill"):
        tp.phase3_ai_metadata_alignment(sorted(blocks), info, None)
    assert SkillInfo.query.count() == 5


def test_visual_enrichment_uses_audited_source_only(section, tmp_path):
    _, info = section
    _, blocks = extract(tmp_path, info)
    tp.phase3_ai_metadata_alignment(sorted(blocks), info, None)
    meta = tp._DOCX_BLOCK_META
    anchors = build_anchors_from_block_meta(meta, info)
    rows = [SimpleNamespace(id=i + 1, source_description=key, problem_text=m["problem_text"],
                            problem_type=m["source_type"], skill_id=m["formal_skill_id"],
                            notes=json.dumps({"question_anchor": anchor}, ensure_ascii=False))
            for i, ((key, m), anchor) in enumerate(zip(meta.items(), anchors))]
    before = [r.notes for r in rows]
    report = enrich_textbook_examples_with_pdf_visuals(
        pdf_path=PDF, examples=rows, curriculum_info=info, project_root=tmp_path,
        debug_dir=tmp_path / "debug", write_notes=False, dpi=72)
    assert report["errors"] == 0
    assert report["questions_matched"] == report["high_confidence"] == 20
    assert report["visual_candidates"] == 6
    assert [r.notes for r in rows] == before
    assert TextbookExample.query.count() == 0
    probe = [{"source_description": "例2", "visual_bbox": None}]
    fake = tmp_path / "different.pdf"
    fake.write_bytes(b"different edition")
    assert correct_pdf_regions(copy.deepcopy(probe), [], fake, info) == probe
    assert correct_pdf_regions(copy.deepcopy(probe), [], PDF, {**info, "section_code": "1-2"}) == probe


def test_profile_does_not_apply_to_other_sources(section):
    _, info = section
    assert is_b2_11(info)
    for change in ({"section_code": "1-2"}, {"volume": "數學B1"},
                   {"source_scope": "chapter_self_assessment"},
                   {"parse_filename": "another.docx"}):
        assert not is_b2_11({**info, **change})


def test_real_pipeline_stops_before_db_write(section, tmp_path):
    app, info = section
    from core.textbook_importer_v3_orchestrate import build_curriculum_info_for_v3_import

    def build(**kwargs):
        return build_curriculum_info_for_v3_import(**kwargs, apply_policy=False)

    with patch.object(pipeline, "_latex_output_path", return_value=tmp_path / "converted.docx"), \
            patch.object(pipeline, "build_curriculum_info_for_v3_import", side_effect=build), \
            patch.object(tp, "phase4_absolute_hydrate_and_save", side_effect=AssertionError("DB write")):
        report = pipeline.run_v3_pair_pipeline(
            project_root=tmp_path, docx_path=SOURCE, pdf_path=PDF,
            curriculum="vocational", volume="數學B2", allow_phase4=False, app=app)
    assert report["ok"], report.get("error")
    assert report["curriculum_info"]["chapter"] == info["chapter"]
    assert report["metrics"]["ai_alignment"]["phase3_questions"] == 20
    assert report["metrics"]["gemini"]["request_count"] == 0
    assert report["metrics"]["db_write"] == {"skipped": True}
    assert SkillInfo.query.count() == 5 and TextbookExample.query.count() == 0
