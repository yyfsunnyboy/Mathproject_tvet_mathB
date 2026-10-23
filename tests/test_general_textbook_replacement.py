import hashlib
import io
import sqlite3
from pathlib import Path
from types import SimpleNamespace

import pytest
from sqlalchemy import Column, Integer, String, create_engine, text
from sqlalchemy.orm import declarative_base, sessionmaker
from werkzeug.datastructures import FileStorage
from werkzeug.datastructures import MultiDict

from core.general_textbook_replacement import (
    backup_sqlite_database, inspect_general_pair, replace_old_rows_in_session,
)

Base = declarative_base()


class Example(Base):
    __tablename__ = "textbook_examples"
    id = Column(Integer, primary_key=True)
    skill_id = Column(String)
    source_curriculum = Column(String)
    source_volume = Column(String)
    source_chapter = Column(String)
    notes = Column(String)


def _session(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'isolated.db'}")
    Base.metadata.create_all(engine)
    with engine.begin() as conn:
        for name, column in (
            ("gencode_component_tracker", "textbook_example_id"),
            ("quiz_attempts", "question_id"),
            ("practice_attempts", "question_uid"),
            ("mistake_notebook_entries", "source_question_uid"),
            ("mistake_logs", "skill_id"),
            ("progress", "skill_id"),
        ):
            columns = '"skill_id" TEXT' if column == "skill_id" else f'"{column}" TEXT, "skill_id" TEXT'
            conn.execute(text(f'CREATE TABLE "{name}" ({columns})'))
    return sessionmaker(bind=engine)(), engine


def _policy(old_ids):
    return {"preview": {"volume": "數學1", "chapter": "單元3 指數", "docx_sha256": "a" * 64}, "old_ids": old_ids}


def _parsed():
    return {"chapters": [{"chapter_title": "3 指數"}]}


def test_first_import_and_scoped_replacement(tmp_path):
    session, engine = _session(tmp_path)
    session.add_all([
        Example(id=1, skill_id="gh_A", source_curriculum="general", source_volume="數學1", source_chapter="單元3 指數"),
        Example(id=2, skill_id="gh_A", source_curriculum="general", source_volume="數學1", source_chapter="單元2 絕對值"),
        Example(id=3, skill_id="vh_A", source_curriculum="vocational", source_volume="數學B1", source_chapter="單元3 指數"),
    ])
    session.commit()
    backup = tmp_path / "backup.db"
    backup_sqlite_database(tmp_path / "isolated.db", backup)
    assert sqlite3.connect(backup).execute("select count(*) from textbook_examples").fetchone()[0] == 3
    assert replace_old_rows_in_session(session, Example, _policy([1]), {"curriculum": "general", "volume": "數學1"}, _parsed()) == 1
    session.add(Example(id=4, skill_id="gh_A", source_curriculum="general", source_volume="數學1", source_chapter="單元3 指數"))
    session.commit()
    assert [(r.id, r.source_curriculum) for r in session.query(Example).order_by(Example.id)] == [(2, "general"), (3, "vocational"), (4, "general")]
    session.close(); engine.dispose()


def test_failed_import_rolls_back_old_rows_and_relation_guard(tmp_path):
    session, engine = _session(tmp_path)
    session.add(Example(id=1, skill_id="gh_A", source_curriculum="general", source_volume="數學1", source_chapter="單元3 指數"))
    session.commit()
    replace_old_rows_in_session(session, Example, _policy([1]), {"curriculum": "general", "volume": "數學1"}, _parsed())
    session.add(Example(id=4, skill_id="gh_A", source_curriculum="general", source_volume="數學1", source_chapter="單元3 指數"))
    session.rollback()
    assert [r.id for r in session.query(Example).all()] == [1]
    session.execute(text("INSERT INTO practice_attempts (skill_id) VALUES ('gh_A')"))
    session.commit()
    with pytest.raises(ValueError, match="relation_blocked"):
        replace_old_rows_in_session(session, Example, _policy([1]), {"curriculum": "general", "volume": "數學1"}, _parsed())
    session.rollback()
    assert [r.id for r in session.query(Example).all()] == [1]
    session.close(); engine.dispose()


def test_scope_change_and_multiple_sources_fail_closed(tmp_path):
    session, engine = _session(tmp_path)
    session.add(Example(id=1, skill_id="gh_A", source_curriculum="general", source_volume="數學1", source_chapter="單元3 指數"))
    session.commit()
    with pytest.raises(ValueError, match="scope_changed"):
        replace_old_rows_in_session(session, Example, _policy([9]), {"curriculum": "general", "volume": "數學1"}, _parsed())
    with pytest.raises(ValueError, match="chapter_mismatch"):
        replace_old_rows_in_session(session, Example, _policy([1]), {"curriculum": "general", "volume": "數學1"}, {"chapters": [{"chapter_title": "4 其他"}]})
    session.rollback()
    session.close(); engine.dispose()


def test_existing_unit3_remains_unchanged():
    root = Path(__file__).resolve().parents[1]
    db = root / "instance/kumon_math.db"
    if not db.is_file():
        pytest.skip("local Unit 3 DB unavailable")
    with sqlite3.connect(f"file:{db.as_posix()}?mode=ro", uri=True) as connection:
        rows = connection.execute("select id,problem_text from textbook_examples where source_curriculum='general' and source_volume='數學1' and source_chapter='單元3 指數' order by id").fetchall()
    assert len(rows) == 38 and [r[0] for r in rows] == list(range(12339, 12377))
    assert "0.7" in dict(rows)[12373] and "1.6" in dict(rows)[12373]


def test_general_pair_preview_reads_unit3_without_importing():
    root = Path(__file__).resolve().parents[1]
    directory = root / "textbook_import/uploaded_sources/general/數學1"
    docx = next(directory.glob("8a1c1765b10a_*.docx"), None)
    pdf = next(directory.glob("ae81389d2b87_*.pdf"), None)
    if not docx or not pdf:
        pytest.skip("local Unit 3 originals unavailable")
    a = FileStorage(stream=io.BytesIO(docx.read_bytes()), filename=docx.name)
    b = FileStorage(stream=io.BytesIO(pdf.read_bytes()), filename=pdf.name)
    preview = inspect_general_pair(a, b, "數學1")
    assert preview["chapter"] == "單元3 指數"
    assert preview["question_count"] == 38
    assert preview["docx_sha256"] == hashlib.sha256(docx.read_bytes()).hexdigest()


def test_web_replacement_previews_then_backs_up_isolated_db(tmp_path, monkeypatch):
    import config as cfg
    import core.routes.admin as admin
    import core.general_textbook_replacement as replacement
    from app import create_app
    from models import db, SkillInfo, TextbookExample

    uri = "sqlite:///" + (tmp_path / "web.db").as_posix()
    monkeypatch.setattr(cfg.Config, "SQLALCHEMY_DATABASE_URI", uri)
    app = create_app()
    app.root_path = str(tmp_path)
    app.secret_key = "isolated-test"
    monkeypatch.setattr(admin, "current_user", SimpleNamespace(is_admin=True, role="teacher"))
    monkeypatch.setattr(admin, "resolve_gemini_api_key", lambda: ("test-key", "test"))
    monkeypatch.setattr(admin, "url_for", lambda *_args, **_kwargs: "/test")
    monkeypatch.setattr(admin, "enqueue_v1_textbook_upload_batch", lambda *_args, **_kwargs: "isolated-task")
    monkeypatch.setattr(replacement, "inspect_general_pair", lambda *_args: {
        "curriculum": "general", "volume": "數學1", "chapter": "單元3 指數",
        "question_count": 2, "docx_sha256": "a" * 64, "pdf_sha256": "b" * 64,
    })
    def upload_data(token=""):
        return MultiDict([
            ("curriculum", "general"), ("volume", "數學1"), ("grade", "10"),
            ("publisher", "longteng"), ("import_mode", "update_existing"),
            ("general_replacement_token", token),
            ("textbook_docx[]", (io.BytesIO(b"docx"), "3.docx")),
            ("textbook_pdf[]", (io.BytesIO(b"pdf"), "3.pdf")),
        ])
    with app.app_context():
        db.create_all()
        db.session.add(SkillInfo(skill_id="gh_Test", skill_en_name="Test", skill_ch_name="Test", description="Test", gemini_prompt="Test"))
        db.session.add(TextbookExample(skill_id="gh_Test", source_curriculum="general", source_volume="數學1",
                                       source_chapter="單元3 指數", source_section="3.指數", source_description="例題1", problem_text="old"))
        db.session.commit()
        with app.test_request_context("/textbook_importer_v3", method="POST", data=upload_data()):
            response, status = admin.admin_textbook_importer_v3.__wrapped__()
        assert status == 409 and response.get_json()["replacement"]["old_count"] == 1
        token = response.get_json()["replacement_token"]
        with app.test_request_context("/textbook_importer_v3", method="POST", data=upload_data(token)):
            response, status = admin.admin_textbook_importer_v3.__wrapped__()
        assert status == 200 and response.get_json()["task_id"] == "isolated-task"
        assert len(list((tmp_path / "backups").glob("general_*_pre_replace.db"))) == 1
        assert TextbookExample.query.count() == 1  # worker was mocked; preview never deletes


def _bindings_session(tmp_path):
    session, engine = _session(tmp_path)
    with engine.begin() as conn:
        conn.execute(text("CREATE TABLE skill_curriculum (id INTEGER PRIMARY KEY, skill_id TEXT, curriculum TEXT, volume TEXT, chapter TEXT)"))
        conn.execute(text("CREATE TABLE skills_info (skill_id TEXT)"))
        for sid in ("gh_New", "gh_Exercises", "vh_X"):
            conn.execute(text("INSERT INTO skills_info VALUES (:s)"), {"s": sid})
        rows = [
            (1, "gh_New", "general", "數學1", "單元3 指數"),
            (2, "gh_Exercises", "general", "數學1", "單元3 指數"),          # stale: label skill, no example in unit
            (3, "gh_Exercises", "general", "數學3A", "單元3 三角"),        # other chapter: must stay
            (4, "gh_New", "general", "數學1", "單元2 絕對值"),            # other chapter: must stay
            (5, "vh_X", "vocational", "數學1", "單元3 指數"),             # other curriculum: must stay
        ]
        for r in rows:
            conn.execute(text("INSERT INTO skill_curriculum VALUES (:a,:b,:c,:d,:e)"), dict(zip("abcde", r)))
    session.add(Example(id=10, skill_id="gh_New", source_curriculum="general", source_volume="數學1", source_chapter="單元3 指數"))
    session.add(Example(id=11, skill_id="gh_Exercises", source_curriculum="general", source_volume="數學3A", source_chapter="單元3 三角"))
    session.commit()
    return session, engine


def test_stale_unit_bindings_removed_idempotently_and_scoped(tmp_path):
    from core.general_textbook_replacement import cleanup_stale_chapter_skill_bindings, stale_chapter_skill_bindings

    session, engine = _bindings_session(tmp_path)
    assert stale_chapter_skill_bindings(session, "數學1", "單元3 指數") == [(2, "gh_Exercises")]
    assert cleanup_stale_chapter_skill_bindings(session, _policy([])) == 1
    assert cleanup_stale_chapter_skill_bindings(session, _policy([])) == 0
    session.commit()
    ids = [r[0] for r in session.execute(text("select id from skill_curriculum order by id"))]
    assert ids == [1, 3, 4, 5]
    assert session.execute(text("select count(*) from skills_info")).scalar() == 3  # skills untouched
    session.close(); engine.dispose()


def test_stale_binding_cleanup_rolls_back_with_failed_import(tmp_path):
    from core.general_textbook_replacement import cleanup_stale_chapter_skill_bindings

    session, engine = _bindings_session(tmp_path)
    cleanup_stale_chapter_skill_bindings(session, _policy([]))
    session.rollback()
    assert session.execute(text("select count(*) from skill_curriculum")).scalar() == 5
    session.close(); engine.dispose()


def test_layout_labels_never_become_skills():
    from core.textbook_processor import is_non_skill_bucket

    for name in ("戊.習題", "戊習題", "3習題", "習題", "隨堂練習", "例題", "自我評量"):
        assert is_non_skill_bucket(name, "Whatever"), name
    for name in ("甲.正整數指數與指數律", "有理數指數與指數律", "乘法公式"):
        assert not is_non_skill_bucket(name, "RationalExponentsAndLaws"), name


def test_section_marker_is_not_part_of_skill_display_name():
    from core.textbook_processor import strip_general_high_section_marker

    assert strip_general_high_section_marker("甲 直線的斜率") == "直線的斜率"
    assert strip_general_high_section_marker("己.二元一次不等式") == "二元一次不等式"
    assert strip_general_high_section_marker("甲. 圓與直線關係的判定") == "圓與直線關係的判定"
    assert strip_general_high_section_marker("常用對數") == "常用對數"
    assert strip_general_high_section_marker("甲烷") == "甲烷"


def test_answer_instruction_is_not_explanation_and_unit4_blocks_all_kept():
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from core.textbook_processor import classify_non_question_block, scan_general_high_layout

    assert classify_non_question_block("試問：太陽的體積為地球的幾倍？（將答案取2位有效數字並以科學記號表示。）") is None
    assert classify_non_question_block("說明：科學記號的定義。") == "concept_explanation"
    assert classify_non_question_block("答案：5") == "concept_explanation"

    source = next(Path("textbook_import/uploaded_sources/general/數學1").glob("*單元04-科學記號與常用對數.docx"))
    doc = Document(str(source))
    chunks = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            value = Paragraph(child, doc).text.strip()
            if value:
                chunks.append(value)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, doc).rows:
                chunks.extend(cell.text.strip() for cell in row.cells if cell.text.strip())
    scan = scan_general_high_layout("\n".join(chunks))
    assert len(scan["items"]) == 26 == len(scan["blocks"])
    assert {t: classify_non_question_block(b) for t, b in scan["blocks"].items() if classify_non_question_block(b)} == {}


def test_fill_in_instruction_is_a_question_and_unit5_blocks_all_kept():
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from core.textbook_processor import classify_non_question_block, scan_general_high_layout

    assert classify_non_question_block("在下列的表格中填入適當的答案：") is None
    assert classify_non_question_block("答案：5") == "concept_explanation"

    source = next(Path("textbook_import/uploaded_sources/general/數學1").glob("*單元05-直線方程式.docx"))
    doc = Document(str(source))
    chunks = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            value = Paragraph(child, doc).text.strip()
            if value:
                chunks.append(value)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, doc).rows:
                chunks.extend(cell.text.strip() for cell in row.cells if cell.text.strip())
    scan = scan_general_high_layout("\n".join(chunks))
    assert len(scan["items"]) == 46 == len(scan["blocks"])
    assert {t: classify_non_question_block(b) for t, b in scan["blocks"].items() if classify_non_question_block(b)} == {}


def _scan_general_source(pattern):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from core.textbook_processor import scan_general_high_layout

    source = next(Path("textbook_import/uploaded_sources/general/數學1").glob(pattern))
    doc = Document(str(source))
    chunks = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            value = Paragraph(child, doc).text.strip()
            if value:
                chunks.append(value)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, doc).rows:
                chunks.extend(cell.text.strip() for cell in row.cells if cell.text.strip())
    return scan_general_high_layout("\n".join(chunks))


def test_instruction_stem_with_explain_word_is_a_question_and_unit6_blocks_all_kept():
    from core.textbook_processor import classify_non_question_block

    stem = "將下列方程式化成\(\left ( { x-h } \right ) ^ { 2 } +\left ( { y-k } \right ) ^ { 2 } =\ell\)的形式，並說明它們的圖形。"
    assert classify_non_question_block(stem) is None
    assert classify_non_question_block("說明：圓的標準式。") == "concept_explanation"
    assert classify_non_question_block("答案：5") == "concept_explanation"

    scan = _scan_general_source("*單元06-圓方程式.docx")
    assert len(scan["items"]) == 32 == len(scan["blocks"])
    assert {t: classify_non_question_block(b) for t, b in scan["blocks"].items() if classify_non_question_block(b)} == {}


def test_source_inventory_is_the_authority_for_question_blocks(monkeypatch):
    """A block whose title the authoritative scan confirmed is never dropped as exposition; unknown titles still are."""
    from core import textbook_processor as T

    exposition_like = "說明下列各圓的圓心與半徑。"  # no instruction lead: the text classifier alone calls it exposition
    assert T.classify_non_question_block(exposition_like) == "concept_explanation"
    monkeypatch.setattr(T, "_DOCX_IMPORT_CONTEXT", {"title_inventory_items": [{"canonical_title": "例題4", "raw_title": "例題4"}]})
    assert T.resolve_question_block_kind(exposition_like, "例題4") is None
    assert T.resolve_question_block_kind(exposition_like, "課文說明") == "concept_explanation"  # not in the inventory
    assert T.resolve_question_block_kind("", "例題4") == "narration"  # empty text is still not a question
    assert T.resolve_question_block_kind("[BLOCK_IMAGE]", "例題4") == "figure_caption"
    monkeypatch.setattr(T, "_DOCX_IMPORT_CONTEXT", {})
    assert T.resolve_question_block_kind(exposition_like, "例題4") == "concept_explanation"  # no inventory -> text only


def test_all_stored_general_sources_keep_every_inventory_block():
    from core.textbook_processor import classify_non_question_block

    counts = {}
    for source in sorted(Path("textbook_import/uploaded_sources/general/數學1").glob("*.docx")):
        scan = _scan_general_source(source.name)
        dropped = {t: classify_non_question_block(b) for t, b in scan["blocks"].items() if classify_non_question_block(b)}
        assert dropped == {}, (source.name, dropped)
        counts[source.name[-20:]] = len(scan["items"])
    assert 32 in counts.values() and 46 in counts.values() and 26 in counts.values()
