# -*- coding: utf-8 -*-
"""Section identity parsing for curriculum binding (B2 Ch4 4-2 hotfix)."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from core.mathb_concept_heading import (
    is_exercise_block_section_title,
    parse_mathb_section_heading,
    section_identities_match,
)
from core.textbook_importer_v3_docx import extract_docx_skill_headings

ROOT = Path(__file__).resolve().parents[1]
B2_4_2_LATEX = next(
    (
        ROOT
        / "textbook_import"
        / "source"
        / "vocational"
        / "math_B2"
    ).glob("*4-2*_Latex.docx"),
    None,
)


@pytest.mark.parametrize(
    "line",
    [
        "4-2 圓與直線的關係",
        "4-2圓與直線的關係",
        "4-2  圓與直線的關係",
        "４－２　圓與直線的關係",
    ],
)
def test_section_heading_accepted_variants(line):
    hit = parse_mathb_section_heading(line, expected_section_code="4-2")
    assert hit is not None
    assert hit["section_code"] == "4-2"
    assert hit["section_title"] == "圓與直線的關係"


@pytest.mark.parametrize(
    "line",
    [
        "4-2.1 點與圓的關係",
        "4-2.1點與圓的關係",
        "4-2 習題",
        "4-2 基礎題",
        "4-2 進階題",
        "4-2 練習題",
        "4-2 自我評量",
        "4-2習題",
    ],
)
def test_section_heading_rejected_non_identity(line):
    assert parse_mathb_section_heading(line, expected_section_code="4-2") is None


def test_exercise_block_title_classifier():
    assert is_exercise_block_section_title("習題")
    assert is_exercise_block_section_title("基礎題")
    assert not is_exercise_block_section_title("圓與直線的關係")


def test_same_section_ignores_whitespace_between_code_and_title():
    assert section_identities_match(
        "4-2圓與直線的關係",
        "4-2 圓與直線的關係",
        expected_section_code="4-2",
    )
    assert not section_identities_match(
        "4-2 習題",
        "4-2 圓與直線的關係",
        expected_section_code="4-2",
    )


def test_extract_prefers_true_section_over_later_exercise(tmp_path):
    """No-space true heading + later ``4-2 習題`` must not flip section identity."""
    path = tmp_path / "section_identity.docx"
    doc = Document()
    doc.add_paragraph("4-2圓與直線的關係")
    doc.add_paragraph("4-2.1 點與圓的關係")
    doc.add_paragraph("內文")
    doc.add_paragraph("4-2 習題")
    doc.add_paragraph("1\t試判斷")
    doc.save(path)

    audit = extract_docx_skill_headings(str(path), section_code="4-2")
    heading = audit["section_heading"]
    assert heading is not None
    assert heading["section_code"] == "4-2"
    assert heading["section_title"] == "圓與直線的關係"
    assert "習題" not in heading["source_heading_text"]
    assert section_identities_match(
        heading,
        "4-2 圓與直線的關係",
        expected_section_code="4-2",
    )


@pytest.mark.skipif(B2_4_2_LATEX is None, reason="B2 4-2 Latex docx not in workspace")
def test_b2_ch4_4_2_latex_section_identity_not_exercise():
    audit = extract_docx_skill_headings(str(B2_4_2_LATEX), section_code="4-2")
    heading = audit["section_heading"]
    assert heading is not None
    assert heading["section_title"] == "圓與直線的關係"
    assert heading["source_heading_text"].replace(" ", "") == "4-2圓與直線的關係"
    codes = [c["concept_code"] for c in audit["skill_candidates"]]
    assert codes == ["4-2.1", "4-2.2", "4-2.3", "4-2.4"]
    assert audit["unresolved_heading_count"] == 0
