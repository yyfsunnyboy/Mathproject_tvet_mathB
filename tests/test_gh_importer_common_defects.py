# -*- coding: utf-8 -*-
"""Regression tests for the general-high importer defects found in 數學2 單元1 (fixed in the shared parser).

Every expected value is what the authoritative DOCX/PDF prints.  Fixtures are synthetic and built in tmp_path;
nothing here opens the production database or writes to the production uploads/ directory.
"""
from __future__ import annotations

import json
import os
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from core.textbook_processor import (
    docx_paragraph_text_faithful,
    extract_converted_latex_docx,
    scan_general_high_layout,
)
from core.textbook_stage1_gate import evaluate_row, evaluate_stage1_gate

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(autouse=True)
def production_assets_untouched():
    uploads = ROOT / "uploads"
    before = sum(len(f) for _, _, f in os.walk(uploads)) if uploads.exists() else 0
    yield
    after = sum(len(f) for _, _, f in os.walk(uploads)) if uploads.exists() else 0
    assert before == after


def _sym(par, char: str, font: str = "Symbol"):
    el = OxmlElement("w:sym")
    el.set(qn("w:font"), font)
    el.set(qn("w:char"), char)
    par.add_run()._r.append(el)


# ---------------------------------------------------------------- symbol / script fidelity (12671, 12674, 12676)
def test_minus_sign_in_example5_is_kept():  # 例題5: a, 2, b-1 成等差
    p = Document().add_paragraph()
    p.add_run("已知a, 2, b三數成等比數列，且a, 2, b ")
    _sym(p, "F02D")
    p.add_run(" 1三數成等差數列，求a的值。")
    text = docx_paragraph_text_faithful(p)
    assert "b − 1三數" in text.replace("  ", " ") or "b − 1" in text
    assert "b  1" not in text


def test_negative_terms_in_practice3_are_kept():  # 隨堂練習3: 3, -6, 12, ..., -384
    p = Document().add_paragraph()
    p.add_run("已知3, ")
    _sym(p, "F02D")
    p.add_run("6, 12,…, ")
    _sym(p, "F02D")
    p.add_run("384構成一個等比數列")
    assert docx_paragraph_text_faithful(p) == "已知3, −6, 12,…, −384構成一個等比數列"


def test_superscript_in_practice5_is_kept():  # 隨堂練習5: 1, r, r^2
    p = Document().add_paragraph()
    p.add_run("已知成等比的1, r, r")
    p.add_run("2").font.superscript = True
    p.add_run("三數恰為另一等差數列之第1, 2, 5項")
    text = docx_paragraph_text_faithful(p)
    assert "\\(r ^ { 2 }\\)" in text and "r2" not in text


def test_script_inside_existing_math_is_not_wrapped_twice():
    p = Document().add_paragraph()
    p.add_run("\\(x")
    p.add_run("2").font.superscript = True
    p.add_run("+1\\)")
    assert docx_paragraph_text_faithful(p) == "\\(x^ { 2 }+1\\)"


def test_unaffected_paragraph_is_returned_unchanged():
    p = Document().add_paragraph("設數列\\(a_{ n }\\)的首項為5。")
    assert docx_paragraph_text_faithful(p) == p.text


def test_unknown_symbol_stays_an_explicit_failure_marker():
    p = Document().add_paragraph()
    _sym(p, "F0FF", "Symbol")
    assert "MATH_PARSE_FAILED:symbol:Symbol:F0FF" in docx_paragraph_text_faithful(p)


def test_table_cell_symbols_are_kept_by_the_converted_docx_reader(tmp_path):
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    cell.add_run("x ")
    _sym(cell, "F0B4")
    cell.add_run(" y")
    path = tmp_path / "t.docx"
    doc.save(path)
    content, _meta = extract_converted_latex_docx(str(path))
    assert "×" in content[1]


# ---------------------------------------------------------------- layout scan: pseudo item, truncation, no over-merge
def _scan(text: str, body: list[str] | None = None):
    return scan_general_high_layout(text, body)


def test_sidebar_sentence_is_not_a_practice():  # 隨堂練習8 偽題
    text = "\n".join([
        "隨堂練習",
        "取一個線段，將其等分成3個線段，重複這樣的步驟，如下圖所示。",
        "依此規律可畫第5圖、第6圖、……，並設\\(a_{ n }\\)為第n圖中的線段個數（\\(a_{ 1 } =1\\),",
        "\\(a_{ 2 } =2\\)）。",
        "(1)寫出數列\\(\\left \\langle { a_{ n } } \\right \\rangle\\)的遞迴關係式。",
        "(2)求\\(a_{ 9 }\\)的值。",
        "隨堂練習中，一直重複其步驟所得到的圖形稱為康托集，它是由德國數學家康托所提出的。",
        "例題8",
        "古代陣法中，士兵依右圖排列。設\\(a_{ n }\\)是前n列士兵的總人數。",
        "(1)求\\(a_{ 1 }\\)的值。",
        "隨堂練習",
        "將所有的正整數依序排列如右圖所示。",
        "(1)寫出遞迴關係式。",
    ])
    scan = _scan(text)
    titles = [i["canonical_title"] for i in scan["items"]]
    assert titles == ["隨堂練習1", "例題8", "隨堂練習2"]
    assert "康托集" not in scan["blocks"]["隨堂練習1"]
    # 練習7 (here 隨堂練習1) keeps the figure lead-in, the sequence definition and both questions
    stem = scan["blocks"]["隨堂練習1"]
    assert "如下圖所示" in stem and "(1)" in stem and stem.rstrip().endswith("\\(a_{ 9 }\\)的值。")


def test_ask_less_stem_continues_over_plain_lines_and_choices():  # 隨堂練習2
    text = "\n".join([
        "隨堂練習",
        "伸出左手，從大拇指開始，按照右圖的方式依序數數字：",
        "1, 2, 3, 4, 5, 6, 7, 8, 9, 10,…。",
        "設\\(a_{ n }\\)是第n次對應到大拇指時所數到的數字。",
        "(1)\t數數看：寫出\\(a_{ 1 } ,\\ a_{ 2 }\\)的值。",
        "(2)\t數到999時，所對應的是哪根手指頭？（單選）",
        "①大拇指　②食指　③中指　④無名指　⑤小指。",
        "(二)\t等比數列",
    ])
    stem = _scan(text)["blocks"]["隨堂練習1"]
    assert "1, 2, 3, 4" in stem and "設\\(a_{ n }\\)" in stem and "(1)" in stem and "⑤小指。" in stem
    assert "等比數列" not in stem


def test_proof_example_keeps_claim_line_and_records_the_proof():  # 例題10
    text = "\n".join([
        "例題10",
        "使用數學歸納法證明：對於所有的正整數n,",
        "\\(a_{ n } =4 ^ { n } +2\\)",
        "恆為6的倍數。",
        "\\(\\o\\ac(■,證)\\)\t①\t當\\(n=1\\)時，\\(a_{ 1 } =4 ^ { 1 } +2=6\\)，得\\(a_{ 1 }\\)是6的倍數。",
        "②\t設\\(n=k\\)時敘述成立。",
    ])
    scan = _scan(text)
    assert scan["blocks"]["例題10"].endswith("恆為6的倍數。")
    assert "當\\(n=1\\)時" in scan["solutions"]["例題10"]


def test_proof_practice_keeps_claim_line():  # 隨堂練習 (9^n+3 恆為12的倍數)
    text = "隨堂練習\n使用數學歸納法證明：對於所有的正整數n,\n\\(9 ^ { n } +3\\)\n恆為12的倍數。\n例題11\n用紅、綠兩色的圖案拼成圖形，如下圖所示。"
    assert _scan(text)["blocks"]["隨堂練習1"].endswith("恆為12的倍數。")


def test_complete_stem_does_not_absorb_following_exposition():
    text = "\n".join([
        "隨堂練習",
        "在同一坐標平面上，\\(y=x ^ { 3 } -4x\\)的圖形是否可以經由適當的平移而完全重疊？若是可以，如何平移呢？",
        "(四)三次函數\\(y=ax ^ { 3 } +bx ^ { 2 } +cx+d\\)的圖形",
        "隨堂練習",
        "解二次不等式\\(9x ^ { 2 } -6x+1\\le 0\\)。",
        "當判別式\\(D < 0\\)時，拋物線與x軸不相交。",
        "底下練習一道綜合等差數列與等比數列的題目。",
        "例題5",
        "已知a, 2, b三數成等比數列，求a的值。",
    ])
    blocks = _scan(text)["blocks"]
    assert "(四)" not in blocks["隨堂練習1"] and "判別式" not in blocks["隨堂練習2"]
    assert "底下練習" not in blocks["隨堂練習2"]


def test_running_text_paragraph_ends_a_stem():
    body = ["上述的幾個例子是我們熟悉的等差或等比數列所產生的遞迴關係式。"]
    text = "隨堂練習\n依此規律可畫第5圖，如下圖所示。\n" + body[0] + "\n例題8\n求\\(a_{ 1 }\\)的值。"
    assert body[0] not in _scan(text, body)["blocks"]["隨堂練習1"]


# ---------------------------------------------------------------- Stage 1 gate: import complete != READY
def _row(text, notes=None, solution="", ptype="in_class_practice"):
    return {"id": 1, "problem_text": text, "detailed_solution": solution, "problem_type": ptype,
            "source_description": "隨堂練習1 [x]", "notes": json.dumps(notes or {})}


def test_gate_blocks_missing_figure_cut_stem_and_symbol_gap():
    assert "required_figure_missing" in evaluate_row(_row("如下圖所示。求\\(a_{ 1 }\\)的值。", {"pdf_visual_status": "VISUAL_REQUIRED_MISSING"}))["blockers"]
    assert "stem_without_terminal_punctuation" in evaluate_row(_row("伸出左手，按照右圖的方式依序數數字："))["blockers"]
    assert "stem_ends_at_figure_lead_in" in evaluate_row(_row("取一個線段，重複這樣的步驟，如下圖所示。"))["blockers"]
    assert "proof_stem_without_claim" in evaluate_row(_row("使用數學歸納法證明：對於所有的正整數n,\n\\(9 ^ { n } +3\\)"))["blockers"]
    assert "stem_lost_symbol_gap" in evaluate_row(_row("已知a, 2, b  1三數成等差數列，求a的值。"))["blockers"]


def test_gate_sends_uncertain_visual_to_review_and_passes_clean_rows():
    uncertain = {"image_assets": [{"path": "p.png", "visual_status": "needs_review", "needs_crop_review": True}]}
    assert evaluate_row(_row("如下圖所示，求\\(a\\)的值。", uncertain))["reviews"] == ["visual_uncertain"]
    ok = {"image_assets": [{"path": "p.png", "visual_status": "accepted"}]}
    assert evaluate_row(_row("如下圖所示，求\\(a\\)的值。", ok)) == {"blockers": [], "reviews": []}


def test_gate_separates_import_complete_from_stage1_ready():
    gate = evaluate_stage1_gate([_row("伸出左手，按照右圖的方式依序數數字："), _row("求\\(x\\)的值。")])
    assert gate["import_complete"] is True and gate["stage1_ready"] is False and gate["status"] == "BLOCKED"
    assert evaluate_stage1_gate([_row("求\\(x\\)的值。")])["stage1_ready"] is True


def test_import_status_surfaces_the_gate():
    from core.general_v1_import_status import GeneralV1ImportStatus

    st = GeneralV1ImportStatus("t")
    st.observe('INFO: [STAGE1 GATE] status=BLOCKED rows=39 blocked=11 needs_review=0 reasons={"required_figure_missing": 11}')
    result = st.snapshot()["result"]
    assert result["stage1_ready"] is False and result["stage1_gate"]["status"] == "BLOCKED"


# ---------------------------------------------------------------- PDF visuals: crop audit and human-verified regions
def test_crop_joining_a_figure_to_a_distant_mascot_goes_to_review():  # 隨堂練習2 誤裁
    from core.textbook_pdf_visual_acceptance import REASON_CROP_DISCONNECTED_REGIONS, crop_extent_reasons

    hand = {"bbox": [420, 450, 500, 520], "area": 5600.0}
    mascot = {"bbox": [440, 640, 500, 720], "area": 4800.0}
    page = {"drawings": [hand], "images": [mascot], "small_drawings": [], "words": []}
    assert crop_extent_reasons(page, [421, 448, 520, 724]) == [REASON_CROP_DISCONNECTED_REGIONS]
    assert crop_extent_reasons(page, [421, 448, 520, 530]) == []


def test_region_overrides_need_reviewer_and_matching_pdf_hash(tmp_path):
    import fitz

    from core.textbook_pdf_visual import (
        apply_visual_region_overrides,
        fingerprint_file,
        load_visual_region_overrides,
    )

    pdf = tmp_path / "unit.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(str(pdf))
    doc.close()
    good = {"pdf_sha256": fingerprint_file(pdf), "verified_by": "reviewer", "verified_on": "2026-10-25",
            "regions": {"例題6": [{"page": 1, "bbox": [10, 10, 50, 50], "why": "figure"}]}}
    regions, warnings = load_visual_region_overrides(pdf, good)
    assert list(regions) == ["例6"] and warnings == []
    assert load_visual_region_overrides(pdf, {**good, "pdf_sha256": "0" * 64})[1] == ["visual_regions_rejected:pdf_hash_mismatch"]
    assert load_visual_region_overrides(pdf, {**good, "verified_by": ""})[1] == ["visual_regions_rejected:reviewer_missing"]
    rows = [{"source_description": "例題6 [source_type=textbook_example]"}, {"source_description": "例題7 [x]"}]
    assert apply_visual_region_overrides(rows, regions) == ["例6"]
    assert rows[0]["should_mount"] is True and rows[0]["match_method"] == "manual_page_region" and "should_mount" not in rows[1]



# --- question_anchor on 例題 rows when an older printing of the chapter is still in the DB (數學2 單元2) ---------------------------

def test_anchor_is_kept_for_every_source_type_with_unique_reading_order():
    from core.textbook_question_anchor import build_anchors_from_block_meta, detect_anchor_id_collisions, normalize_question_label

    titles = ["例題1", "隨堂練習1", "例題2", "隨堂練習2", "2習題 觀念澄清1", "2習題 基礎題1", "2習題 進階題9"]
    meta = {t: {"anchor": t, "problem_text": f"stem {t}", "source_type": ""} for t in titles}
    anchors = build_anchors_from_block_meta(meta, {"curriculum": "general", "publisher": "longteng", "volume": "數學2", "chapter_index": 2})
    assert [a["source_order"] for a in anchors] == list(range(1, len(titles) + 1))  # reading order, examples included
    assert detect_anchor_id_collisions(anchors) == [] and len({a["anchor_id"] for a in anchors}) == len(titles)
    assert {normalize_question_label(a["question_label"]) for a in anchors} >= {"例1", "例2", "隨堂練習1"}
    assert [a["question_type"] for a in anchors][:2] == ["example", "in_class_practice"]


def test_stale_older_printing_does_not_leave_example_rows_unanchored():
    from types import SimpleNamespace

    from core.textbook_processor import pick_current_import_row

    fresh = SimpleNamespace(id=12707, notes=json.dumps({"source_provenance": {"primary": {"sha256": "a" * 64}}}))
    stale = SimpleNamespace(id=1334, notes=json.dumps({}))          # earlier import of the same chapter: no provenance
    stale_null = SimpleNamespace(id=1335, notes=None)
    assert pick_current_import_row([stale, fresh, stale_null]) is fresh
    # two rows that both carry provenance (two imports of the same source) stay ambiguous: never guess
    other = SimpleNamespace(id=13000, notes=json.dumps({"source_provenance": {"primary": {"sha256": "b" * 64}}}))
    assert pick_current_import_row([fresh, other]) is None
    assert pick_current_import_row([stale, stale_null]) is None
