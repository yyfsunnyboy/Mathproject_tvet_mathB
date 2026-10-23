# -*- coding: utf-8 -*-
"""Regression: existing V1 普高 (curriculum=general) textbook import path."""
from __future__ import annotations

import io
import queue as queue_mod
import uuid

import pytest

from core import textbook_processor as tp
from core.textbook_solution_boundary import split_student_stem_at_solution_marker

# Layout markers only; the content is a neutral placeholder unit.
GENERAL_HIGH_TEXT = "\n".join(
    [
        "3函數",
        "甲 函數的意義",
        "說明文字，介紹函數。",
        "例題1",
        "已知 \\(f(x)=x+1\\)，求 \\(f(2)\\)。",
        "\\(\\o\\ac(■,解)\\)\t\\(f(2)=3\\)。",
        "隨堂練習",
        "已知 \\(f(x)=2x\\)，求 \\(f(3)\\)。",
        "接著說明下一個概念，這是課文敘述。",
        "乙 函數的圖形",
        "(一)線性函數",
        "例題2",
        "畫出 \\(y=x\\) 的圖形。",
        "\\(\\o\\ac(■,解)\\)\t圖形為一直線。",
        "隨堂練習",
        "畫出 \\(y=2x\\) 的圖形。",
        "3習題",
        "觀念澄清",
        "1.下列敘述對的打「○」，錯的打「×」。",
        "(1)\\(f(x)=1\\) 是函數。",
        "一、基礎題",
        "1.求 \\(f(1)\\)。",
        "2.求 \\(f(0)\\)。",
        "二、進階題",
        "3.求 \\(f(-1)\\)。",
    ]
)


def test_scan_general_high_layout_sections_subsections_and_titles():
    scan = tp.scan_general_high_layout(GENERAL_HIGH_TEXT)
    assert scan["chapter"] == {"unit": "3", "title": "函數", "line": "3函數"}
    by_title = {it["canonical_title"]: it for it in scan["items"]}
    assert list(by_title) == [
        "例題1",
        "隨堂練習1",
        "例題2",
        "隨堂練習2",
        "3習題 觀念澄清1",
        "3習題 基礎題1",
        "3習題 基礎題2",
        "3習題 進階題3",
    ]
    assert by_title["例題1"]["section_heading"] == "甲 函數的意義"
    assert by_title["例題1"]["subsection_heading"] == ""
    assert by_title["例題2"]["section_heading"] == "乙 函數的圖形"
    assert by_title["例題2"]["subsection_heading"] == "(一)線性函數"
    assert by_title["3習題 基礎題1"]["kind"] == "chapter_exercise"
    assert by_title["3習題 基礎題1"]["zone"] == "基礎題"
    assert by_title["3習題 觀念澄清1"]["zone"] == "觀念澄清"


def test_scan_general_high_layout_splits_solution_and_stops_practice_at_narrative():
    scan = tp.scan_general_high_layout(GENERAL_HIGH_TEXT)
    assert "解" not in scan["blocks"]["例題1"]
    assert scan["solutions"]["例題1"].startswith("\\(f(2)=3\\)")
    assert "課文敘述" not in scan["blocks"]["隨堂練習1"]
    assert "(1)" in scan["blocks"]["3習題 觀念澄清1"]


UNIT2_LIKE_TEXT = "\n".join(
    [
        "2絕對值",
        "甲 數線上的分點公式",
        "例題8",
        "人體體溫超過38度稱為發燒。",
        "\\(\\o\\ac(■,解)\\)\t依題意可知 \\(x>38\\)。",
        "隨堂練習",
        "世界衛生組織規範男性標準體重的公式為",
        "\\(標準體重=(身高-80)\\times 0.7\\)。",
        "根據此規範回答下列各題：",
        "(1)身高180公分的男性，其標準體重為多少公斤？",
        "(2)已知 \\(|x-a|\\le b\\)，求a, b的值。",
        "例題8 的解法，除了利用代數解法直接求解外，也可以透過幾何的解法求解：",
        "2習題",
        "觀念澄清",
        "下列敘述對的打「○」，錯的打「×」。",
        "(1)數線上 \\(A(a)\\) 與 \\(B(-4)\\) 的距離為 \\(a+4\\)。",
        "(2)若 \\(a<b\\)，則 \\(\\frac{2a+b}{3}>\\frac{a+2b}{3}\\)。",
        "(3)\\(|x-1|\\le 3\\) 與 \\(|1-x|\\le 3\\) 的解相同。",
        "一、基礎題",
        "1.設數線上兩點 \\(A(5)\\)、\\(B(12)\\)。",
    ]
)


def test_example_back_reference_is_not_a_new_example():
    scan = tp.scan_general_high_layout(UNIT2_LIKE_TEXT)
    titles = [it["canonical_title"] for it in scan["items"]]
    assert titles.count("例題8") == 1
    assert "的解法" not in scan["blocks"]["例題8"]


UNIT7_BACK_REFERENCE_TEXT = "\n".join([
    "7圓與直線的關係",
    "甲 圓與直線關係的判定",
    "例題1",
    "判斷圓 \\(C:x^{2}+y^{2}=5\\) 和直線 \\(L:x-y+1=0\\) 的相交情形。",
    "\\(\\o\\ac(■,解)\\)\t解聯立方程式。",
    "隨堂練習",
    "判斷圓 \\(C\\) 與下列直線的相交情形：",
    "(1)\\(L_{1}:x+y-5=0\\)。　(2)\\(L_{2}:x+y-6=0\\)。",
    "例題1解方程式的過程中，利用代入消去法可消去其中一個未知數，",
    "(1)\t當 \\(D>0\\) 時，方程式有兩相異實數解。",
    "例題2",
    "設圓 \\(C\\) 與直線 \\(L:3x+4y+k=0\\) 相切，求k的值。",
])


def test_example_back_reference_sentence_without_particle_is_not_a_new_example():
    scan = tp.scan_general_high_layout(UNIT7_BACK_REFERENCE_TEXT)
    titles = [it["canonical_title"] for it in scan["items"]]
    assert titles == ["例題1", "隨堂練習1", "例題2"]
    assert "判斷圓" in scan["blocks"]["例題1"]
    assert "代入消去法" not in scan["blocks"]["例題1"]
    assert "代入消去法" not in scan["blocks"]["隨堂練習1"]


def test_repeated_bare_example_heading_still_reaches_inventory_gate():
    text = UNIT7_BACK_REFERENCE_TEXT.replace("例題1解方程式的過程中，利用代入消去法可消去其中一個未知數，", "例題1")
    titles = [it["canonical_title"] for it in tp.scan_general_high_layout(text)["items"]]
    assert titles.count("例題1") == 2


def test_practice_lead_in_before_subquestions_stays_in_stem():
    stem = tp.scan_general_high_layout(UNIT2_LIKE_TEXT)["blocks"]["隨堂練習1"]
    assert "根據此規範回答下列各題" in stem
    assert "(1)身高180公分" in stem and "(2)已知" in stem
    assert "的解法" not in stem


def test_exposition_after_practice_subquestions_is_not_a_lead_in():
    text = "\n".join([
        "隨堂練習",
        "將下列各小數化成最簡分數：",
        "(1) 0.732。　　(2) 0.58。",
        "仿照例題2的方式，每個有限小數都可以化成有理數的形式。事實上，",
        "(1)任意有理數都可化成整數、有限小數或循環小數。",
        "(2)整數、有限小數與循環小數都是有理數。",
    ])
    stem = tp.scan_general_high_layout(text)["blocks"]["隨堂練習1"]
    assert stem.endswith("(2) 0.58。")
    assert "仿照" not in stem and "任意有理數" not in stem


@pytest.mark.parametrize(
    "pages, expected",
    [
        ([7, None, None, 7], {1: 7, 2: 7}),
        ([7, None, 8], {}),
        ([None, 3], {}),
        ([3, None], {}),
        ([2, 3], {}),
    ],
)
def test_unmatched_pdf_page_inferred_only_between_same_page_neighbors(pages, expected):
    assert tp.interpolate_unmatched_pdf_pages(pages) == expected


def test_unnumbered_concept_check_items_share_instruction():
    scan = tp.scan_general_high_layout(UNIT2_LIKE_TEXT)
    titles = [it["canonical_title"] for it in scan["items"]]
    assert titles == [
        "例題8", "隨堂練習1",
        "2習題 觀念澄清1", "2習題 觀念澄清2", "2習題 觀念澄清3",
        "2習題 基礎題1",
    ]
    first = scan["blocks"]["2習題 觀念澄清1"]
    assert first.startswith("下列敘述對的打「○」，錯的打「×」。")
    assert "B(-4)" in first and "a<b" not in first
    assert "|1-x|" in scan["blocks"]["2習題 觀念澄清3"]


@pytest.mark.parametrize("line", ["(　)1.數線上兩點的距離。", "□ 1 數線上兩點的距離。", "（ ）(1)數線上兩點的距離。"])
def test_boxed_concept_check_item_starts_item(line):
    text = "\n".join(["2習題", "觀念澄清", line, "一、基礎題", "1.求值。"])
    scan = tp.scan_general_high_layout(text)
    assert [it["canonical_title"] for it in scan["items"]] == ["2習題 觀念澄清1", "2習題 基礎題1"]
    assert scan["blocks"]["2習題 觀念澄清1"].endswith("數線上兩點的距離。")


@pytest.mark.parametrize(
    "text, needs_image",
    [
        ("設數線上兩點 \\(A\\left ( { 5 } \\right )\\)、\\(B\\left ( { 12 } \\right )\\)。求AB的長。", False),
        ("設 A(12), B(-6) 為數線上兩點，求P的坐標。", False),
        ("請寫出數線上 P 點的座標。", True),
        ("如右圖，數線上 A(3)，求 P 點坐標。", True),
    ],
)
def test_number_line_points_given_in_text_need_no_figure(text, needs_image):
    assert tp.classify_image_dependency(text)["needs_image"] is needs_image


def test_source_file_provenance_records_original_name_and_sha256(tmp_path):
    import hashlib

    src = tmp_path / "02-.docx"
    src.write_bytes(b"docx-bytes")
    prov = tp.source_file_provenance(str(src), "02-絕對值-word.docx")
    assert prov == {
        "filename": "02-.docx",
        "original_filename": "02-絕對值-word.docx",
        "size_bytes": 10,
        "sha256": hashlib.sha256(b"docx-bytes").hexdigest(),
    }
    assert tp.source_file_provenance(str(tmp_path / "missing.docx")) is None


SOLUTION_THEN_PROSE_TEXT = "\n".join(
    [
        "3函數",
        "甲 函數的意義",
        "例題1",
        "已知 \\(f(x)=x+1\\)。",
        "說明條件的題幹段落。",
        "求 \\(f(2)\\)。",
        "\\(\\o\\ac(■,解)\\)\t\\(f(2)=3\\)。",
        "故所求為 3。",
        "由上例可知，這是接續的課文敘述。",
        "課文第二段。",
        "例題2",
        "求 \\(f(0)\\)。",
        "\\(\\o\\ac(■,解)\\)\t\\(f(0)=1\\)。",
    ]
)
SOLUTION_THEN_PROSE_BODY = ["說明條件的題幹段落。", "由上例可知，這是接續的課文敘述。", "課文第二段。"]


def test_example_solution_stops_at_running_text_paragraph():
    scan = tp.scan_general_high_layout(SOLUTION_THEN_PROSE_TEXT, body_paragraphs=SOLUTION_THEN_PROSE_BODY)
    assert scan["solutions"]["例題1"] == "\\(f(2)=3\\)。\n故所求為 3。"
    assert "說明條件的題幹段落。" in scan["blocks"]["例題1"]
    assert scan["solutions"]["例題2"] == "\\(f(0)=1\\)。"


def test_example_solution_boundary_unchanged_without_layout_evidence():
    scan = tp.scan_general_high_layout(SOLUTION_THEN_PROSE_TEXT)
    assert scan["solutions"]["例題1"].endswith("課文第二段。")


def test_extract_converted_docx_reports_running_text_paragraphs(tmp_path):
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    prose = doc.add_paragraph("課文敘述段落。")
    prose.paragraph_format.first_line_indent = Pt(24)
    solution = doc.add_paragraph("詳解延續行。")
    solution.paragraph_format.left_indent = Pt(24)
    solution.paragraph_format.first_line_indent = Pt(-12)
    doc.add_paragraph("無縮排標題")
    cell = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    cell.text = "表格內題幹"
    cell.paragraph_format.first_line_indent = Pt(24)
    path = tmp_path / "layout.docx"
    doc.save(path)

    pages, meta = tp.extract_converted_latex_docx(str(path))
    assert meta["body_paragraphs"] == ["課文敘述段落。"]
    assert "表格內題幹" in pages[1]


def test_filled_square_eq_solution_marker_is_recognized():
    stem, solution = split_student_stem_at_solution_marker("題幹\n\\(\\o\\ac(■,解)\\)\t詳解內容")
    assert stem == "題幹"
    assert solution == "詳解內容"


def test_default_layout_keeps_legacy_inventory_scan():
    text = "1-1 數線\n例題1\n求值。\n隨堂練習1\n再求值。"
    assert tp.scan_docx_title_inventory(text) == tp.scan_docx_title_inventory(text, layout="")
    assert all("section_heading" not in it for it in tp.scan_docx_title_inventory(text))


@pytest.mark.parametrize(
    "raw, expected",
    [
        ("第3章 函數", ("單元3 函數", "3.函數")),
        ("單元3 函數", ("單元3 函數", "3.函數")),
        ("3 函數", ("單元3 函數", "3.函數")),
        ("3.函數", ("單元3 函數", "3.函數")),
        ("函數", ("函數", "")),
    ],
)
def test_general_high_chapter_section_naming(raw, expected):
    assert tp.normalize_general_high_chapter_section(raw) == expected


def test_source_verbatim_solution_is_not_trimmed():
    long_solution = "(1)第一段\n\n(2)" + "x" * 600
    assert tp.sanitize_detailed_solution_text(long_solution, source_verbatim=True) == long_solution
    legacy = tp.sanitize_detailed_solution_text(long_solution)
    assert len(legacy) == 500 and not legacy.startswith("(1)")


def test_concept_zone_exercise_is_chapter_exercise():
    assert tp.normalize_source_type_by_title({"title": "3習題 觀念澄清1"}) == "chapter_exercise"
    assert tp.normalize_source_type_by_title({"title": "3習題 基礎題1"}) == "basic_exercise"


def test_process_textbook_file_sets_general_layout_and_mathtype_gate(monkeypatch):
    from app import app

    seen = []

    def _fake_extract(file_path, queue, max_pages=None, import_policy=None):
        seen.append(dict(import_policy or {}))
        return None

    monkeypatch.setattr(tp, "extract_content_from_file", _fake_extract)
    with app.app_context():
        tp.process_textbook_file("x.docx", {"curriculum": "general"}, queue_mod.Queue(), import_policy={})
        tp.process_textbook_file("x.docx", {"curriculum": "vocational"}, queue_mod.Queue(), import_policy={})
    assert seen[0]["docx_layout"] == tp.GENERAL_HIGH_LAYOUT
    assert seen[0]["mathtype_ole_autoconvert"] is True
    assert "docx_layout" not in seen[1]
    assert "mathtype_ole_autoconvert" not in seen[1]


@pytest.fixture()
def admin_app(tmp_path):
    import config as _cfg
    from app import create_app
    from models import User, db

    prev_uri = _cfg.Config.SQLALCHEMY_DATABASE_URI
    _cfg.Config.SQLALCHEMY_DATABASE_URI = "sqlite:///" + str(tmp_path / f"gh_{uuid.uuid4().hex[:8]}.db").replace("\\", "/")
    try:
        app = create_app()
        app.config.update(TESTING=True)
        app.root_path = str(tmp_path)
        with app.app_context():
            admin = User(username=f"admin_{uuid.uuid4().hex[:6]}", password_hash="x", role="admin")
            db.session.add(admin)
            db.session.commit()
            yield app, admin.id
    finally:
        _cfg.Config.SQLALCHEMY_DATABASE_URI = prev_uri


def test_v1_route_dispatches_general_curriculum(admin_app, monkeypatch):
    import core.routes.admin as admin_mod

    app, admin_id = admin_app
    started = []

    class _FakeThread:
        def __init__(self, target=None, args=(), kwargs=None):
            started.append((target, args, kwargs or {}))

        def start(self):
            pass

    monkeypatch.setattr(admin_mod, "resolve_gemini_api_key", lambda: ("k", "test"))
    monkeypatch.setattr(admin_mod.threading, "Thread", _FakeThread)
    client = app.test_client()
    with client.session_transaction() as sess:
        sess["_user_id"] = str(admin_id)
        sess["_fresh"] = True
    resp = client.post(
        "/textbook_importer",
        data={
            "curriculum": "general",
            "publisher": "longteng",
            "grade": "10",
            "volume": "數學1",
            "textbook_pdf": (io.BytesIO(b"PK"), "unit.docx"),
        },
        content_type="multipart/form-data",
    )
    assert resp.status_code == 302
    assert len(started) == 1
    target, args, kwargs = started[0]
    assert target is admin_mod.background_processing
    assert args[3]["curriculum"] == "general"
    assert args[3]["volume"] == "數學1"
    assert kwargs["import_policy"]["docx_formula_source_mode"] == "converted_docx_latex"


def test_background_processing_pairs_word_and_pdf_variants(monkeypatch):
    import core.routes.admin as admin_mod
    from app import app

    calls = []

    def _fake_process(file_path, **kwargs):
        calls.append((file_path, kwargs.get("optional_enrich_pdf_path")))
        return {"status": "success"}

    monkeypatch.setattr(admin_mod.textbook_processor, "process_textbook_file", _fake_process)
    paths = ["u/01--word.docx", "u/01--pdf.pdf", "u/B2_1-1.docx", "u/B2_1-1.pdf"]
    admin_mod.background_processing(paths, queue_mod.Queue(), app.app_context(), {"curriculum": "general"}, True)
    assert ("u/01--word.docx", "u/01--pdf.pdf") in calls
    assert ("u/B2_1-1.docx", "u/B2_1-1.pdf") in calls
    assert len(calls) == 2


@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        ("單元01-實數-課本word檔.docx", "單元01-實數-課本"),
        ("單元01-實數-課本pdf檔.pdf", "單元01-實數-課本"),
        ("單元01-實數-word檔.docx", "單元01-實數"),
        ("單元01-實數-PDF.pdf", "單元01-實數"),
        ("單元01-實數.docx", "單元01-實數"),
        ("word檔中段-實數.docx", "word檔中段-實數"),
    ],
)
def test_general_source_pair_stem_strips_only_trailing_file_type(filename, expected):
    import core.routes.admin as admin_mod

    assert admin_mod.general_source_pair_stem(filename) == expected


def _grouped_calls(monkeypatch, paths, curriculum):
    import core.routes.admin as admin_mod
    from app import app

    calls = []
    monkeypatch.setattr(
        admin_mod.textbook_processor,
        "process_textbook_file",
        lambda file_path, **kwargs: calls.append((file_path, kwargs.get("optional_enrich_pdf_path"))) or {"status": "success"},
    )
    admin_mod.background_processing(paths, queue_mod.Queue(), app.app_context(), {"curriculum": curriculum}, True)
    return calls


@pytest.mark.parametrize(
    ("docx", "pdf"),
    [
        ("u/單元01-實數-課本word檔.docx", "u/單元01-實數-課本pdf檔.pdf"),
        ("u/單元01-實數-word檔.docx", "u/單元01-實數-pdf檔.pdf"),
        ("u/單元01-實數.docx", "u/單元01-實數.pdf"),
    ],
)
def test_general_pairs_file_type_suffixed_sources(monkeypatch, docx, pdf):
    assert _grouped_calls(monkeypatch, [docx, pdf], "general") == [(docx, pdf)]


def test_general_does_not_pair_different_units(monkeypatch):
    calls = _grouped_calls(monkeypatch, ["u/單元01-實數-word檔.docx", "u/單元02-指數-pdf檔.pdf"], "general")
    assert ("u/單元01-實數-word檔.docx", None) in calls
    assert all(pdf is None for _docx, pdf in calls)


def test_vocational_pairing_ignores_general_suffix_normalization(monkeypatch):
    calls = _grouped_calls(monkeypatch, ["u/1-1-課本word檔.docx", "u/1-1-課本pdf檔.pdf"], "vocational")
    assert ("u/1-1-課本word檔.docx", None) in calls
    assert all(pdf is None for _docx, pdf in calls)


def test_v3_preview_receives_general_pair_pattern(admin_app, monkeypatch):
    import core.routes.admin as admin_mod

    app, admin_id = admin_app
    monkeypatch.setattr(admin_mod, "resolve_gemini_api_key", lambda: ("k", "test"))
    client = app.test_client()
    with client.session_transaction() as sess:
        sess["_user_id"] = str(admin_id)
        sess["_fresh"] = True
    html = client.get("/textbook_importer_v3").get_data(as_text=True)
    assert "GENERAL_SOURCE_PAIR_SUFFIX_RE = new RegExp(" in html
    assert "word\\\\s*\\u6a94" in html or "word\\\\s*檔" in html


def test_v3_preview_fallback_pattern_matches_backend():
    from pathlib import Path

    import core.routes.admin as admin_mod

    template = (Path(__file__).resolve().parents[1] / "templates" / "textbook_importer_v3.html").read_text(encoding="utf-8")
    assert "String.raw`" + admin_mod.GENERAL_SOURCE_PAIR_SUFFIX_PATTERN + "`" in template


def test_background_processing_forwards_general_curriculum_to_v1_importer(monkeypatch):
    import core.routes.admin as admin_mod
    from app import app

    received = []
    monkeypatch.setattr(
        admin_mod.textbook_processor,
        "process_textbook_file",
        lambda _file_path, **kwargs: received.append(kwargs["curriculum_info"]) or {"status": "success"},
    )

    admin_mod.background_processing(
        ["u/1-1.docx", "u/1-1.pdf"],
        queue_mod.Queue(),
        app.app_context(),
        {"curriculum": "general", "publisher": "longteng", "grade": "10", "volume": "數學1"},
        False,
    )

    assert received == [{"curriculum": "general", "publisher": "longteng", "grade": "10", "volume": "數學1"}]


def test_background_processing_skips_unpaired_pdf_only_for_general(monkeypatch):
    import core.routes.admin as admin_mod
    from app import app

    calls = []
    monkeypatch.setattr(
        admin_mod.textbook_processor,
        "process_textbook_file",
        lambda file_path, **kwargs: calls.append((file_path, kwargs.get("optional_enrich_pdf_path"))) or {"status": "success"},
    )
    q = queue_mod.Queue()
    admin_mod.background_processing(["u/02--word.docx", "u/lone.pdf"], q, app.app_context(), {"curriculum": "general"}, True)
    assert calls == [("u/02--word.docx", None)]
    assert any("unpaired PDF skipped" in str(m) for m in list(q.queue))

    calls.clear()
    admin_mod.background_processing(["u/lone.pdf"], queue_mod.Queue(), app.app_context(), {"curriculum": "vocational"}, True)
    assert calls == [("u/lone.pdf", None)]


def _run_converted_process(monkeypatch, curriculum, pdf_path):
    from app import app

    enrich_calls = []
    monkeypatch.setattr(tp, "extract_content_from_file", lambda *a, **k: {1: GENERAL_HIGH_TEXT})
    monkeypatch.setattr(tp, "_DOCX_IMPORT_CONTEXT", {"docx_formula_source_mode": "converted_docx_latex"})
    monkeypatch.setattr(
        tp,
        "call_gemini_for_analysis",
        lambda *a, **k: '{"chapters": [{"chapter_title": "單元3 函數", "sections": []}]}',
    )
    monkeypatch.setattr(tp, "write_title_inventory_report", lambda *a, **k: None)
    monkeypatch.setattr(tp, "save_to_database", lambda *a, **k: {"status": "success"})
    monkeypatch.setattr(
        tp,
        "enrich_general_high_pdf_visuals",
        lambda **kwargs: enrich_calls.append(kwargs) or {"ok": True, "status_counts": {}},
    )
    q = queue_mod.Queue()
    with app.app_context():
        result = tp.process_textbook_file(
            "u/03--word.docx",
            {"curriculum": curriculum, "volume": "數學1"},
            q,
            skip_code_gen=True,
            import_policy={"visual_asset_root": "tmp-assets"},
            optional_enrich_pdf_path=pdf_path,
        )
    return result, enrich_calls, [str(m) for m in list(q.queue)]


def test_general_high_paired_pdf_invokes_visual_enrichment(monkeypatch):
    result, calls, _messages = _run_converted_process(monkeypatch, "general", "u/03--pdf.pdf")
    assert result["status"] == "success"
    assert len(calls) == 1
    assert calls[0]["pdf_path"] == "u/03--pdf.pdf"
    assert calls[0]["scan"]["chapter"]["unit"] == "3"
    assert calls[0]["import_policy"]["visual_asset_root"] == "tmp-assets"


def test_general_high_without_pdf_falls_back_without_enrichment(monkeypatch):
    result, calls, messages = _run_converted_process(monkeypatch, "general", None)
    assert result["status"] == "success"
    assert calls == []
    assert "INFO: [PDF VISUAL] skipped reason=paired_pdf_missing" in messages


def test_vocational_import_never_runs_general_high_visual_enrichment(monkeypatch):
    result, calls, messages = _run_converted_process(monkeypatch, "vocational", "u/03--pdf.pdf")
    assert result["status"] == "success"
    assert calls == []
    assert not any("[PDF VISUAL]" in m for m in messages)


def _seed_general_rows(db, rows):
    import json

    from models import SkillInfo, TextbookExample

    db.session.add(SkillInfo(skill_id="gh_visual_test", skill_en_name="t", skill_ch_name="t", description="t", gemini_prompt="t"))
    created = {}
    for title, problem, notes in rows:
        te = TextbookExample(
            skill_id="gh_visual_test",
            source_curriculum="general",
            source_volume="數學1",
            source_chapter="單元3 函數",
            source_section="3.函數",
            source_description=f"{title} [t]",
            problem_text=problem,
            notes=json.dumps(notes, ensure_ascii=False),
        )
        db.session.add(te)
        created[title] = te
    db.session.commit()
    return created


_UNIT2_MAP = [{
    "id": "t_unit2",
    "curriculum": "general",
    "volume": "數學1",
    "chapter": "單元2 絕對值",
    "titles": {"例題1": "gh_PointDivisionFormulaOnNumberLine", "2習題 進階題7": "gh_LinearEquationsWithAbsoluteValue"},
}]


def _seed_unit_rows(db, chapter, rows, skills):
    from models import SkillCurriculum, SkillInfo, TextbookExample

    for sid in skills:
        if not db.session.get(SkillInfo, sid):
            db.session.add(SkillInfo(skill_id=sid, skill_en_name=sid, skill_ch_name=sid, description="t", gemini_prompt="t"))
            db.session.add(SkillCurriculum(skill_id=sid, curriculum="general", grade=10, volume="數學1", chapter=chapter, section="s",
                                           paragraph=f"乙.{sid}_para"))
    created = {}
    for title, skill in rows:
        te = TextbookExample(skill_id=skill, source_curriculum="general", source_volume="數學1", source_chapter=chapter,
                             source_section="s", source_paragraph=f"{skill}_para", source_description=f"{title} [t]",
                             problem_text="p", notes="{}")
        db.session.add(te)
        created[title] = te
    db.session.commit()
    return {t: te.id for t, te in created.items()}


def test_source_skill_map_overrides_ai_skill_only_in_mapped_unit(admin_app):
    import json

    from models import TextbookExample, db

    eq, ineq, pd = "gh_LinearEquationsWithAbsoluteValue", "gh_LinearInequalitiesWithAbsoluteValue", "gh_PointDivisionFormulaOnNumberLine"
    ids = _seed_unit_rows(db, "單元2 絕對值", [("例題1", pd), ("2習題 進階題7", ineq), ("例題9", ineq)], [eq, ineq, pd])
    other = _seed_unit_rows(db, "單元1 實數", [("2習題 進階題7", "gh_other")], ["gh_other"])

    summary = tp.apply_general_high_source_skill_map(
        parsed_data={"chapters": [{"chapter_title": "單元2 絕對值"}]},
        scan={"chapter": {}},
        curriculum_info={"curriculum": "general", "volume": "數學1"},
        maps=_UNIT2_MAP,
    )

    assert summary["applied"] is True
    assert summary["changed"] == [{"title": "2習題 進階題7", "from": ineq, "to": eq}]
    assert summary["unmapped"] == ["例題9"]
    row7 = db.session.get(TextbookExample, ids["2習題 進階題7"])
    assert row7.skill_id == eq
    assert row7.source_paragraph == f"{eq}_para"
    assert json.loads(row7.notes)["skill_binding"] == {
        "source": "general_high_source_skill_map", "map_id": "t_unit2",
        "ai_skill_id": ineq, "ai_source_paragraph": f"{ineq}_para",
    }
    assert db.session.get(TextbookExample, ids["例題1"]).source_paragraph == f"{pd}_para"
    assert json.loads(db.session.get(TextbookExample, ids["例題9"]).notes)["needs_review"] is True
    assert db.session.get(TextbookExample, other["2習題 進階題7"]).skill_id == "gh_other"


def test_source_skill_map_rejects_skill_not_bound_to_chapter(admin_app):
    from models import TextbookExample, db

    ids = _seed_unit_rows(db, "單元2 絕對值", [("例題1", "gh_x")], ["gh_x"])
    bad_map = [dict(_UNIT2_MAP[0], titles={"例題1": "gh_PointDivisionFormulaOnNumberLine"})]
    summary = tp.apply_general_high_source_skill_map(
        parsed_data={"chapters": [{"chapter_title": "單元2 絕對值"}]},
        scan={"chapter": {}},
        curriculum_info={"curriculum": "general", "volume": "數學1"},
        maps=bad_map,
    )
    assert summary["applied"] is False and summary["error"].startswith("skills_not_in_chapter")
    assert db.session.get(TextbookExample, ids["例題1"]).skill_id == "gh_x"


@pytest.mark.parametrize(
    "curriculum, volume, chapter",
    [("vocational", "數學1", "單元2 絕對值"), ("general", "數學2", "單元2 絕對值"), ("general", "數學1", "單元1 實數")],
)
def test_source_skill_map_scope_is_exact(curriculum, volume, chapter):
    assert tp.find_general_high_source_skill_map(curriculum, volume, chapter, _UNIT2_MAP) is None


def test_repository_unit2_source_skill_map_counts():
    entry = tp.find_general_high_source_skill_map("general", "數學1", "單元2 絕對值")
    counts = {}
    for sid in entry["titles"].values():
        counts[sid] = counts.get(sid, 0) + 1
    assert len(entry["titles"]) == 28
    assert counts == {
        "gh_PointDivisionFormulaOnNumberLine": 8,
        "gh_LinearEquationsWithAbsoluteValue": 6,
        "gh_LinearInequalitiesWithAbsoluteValue": 14,
    }


_UNIT7 = "單元7 圓與直線的關係"
_JUDGE, _TANGENT = "gh_JudgingTheRelationshipOfCircleAndLine", "gh_TangentLinesOfACircle"
_UNIT7_TEACHER_FIXED = {"7習題 進階題10", "7習題 進階題12"}


def test_repository_unit7_source_skill_map_counts():
    entry = tp.find_general_high_source_skill_map("general", "數學1", _UNIT7)
    counts = {}
    for sid in entry["titles"].values():
        counts[sid] = counts.get(sid, 0) + 1
    assert len(entry["titles"]) == 36
    assert counts == {_JUDGE: 21, _TANGENT: 15}
    assert all(entry["titles"][t] == _JUDGE for t in _UNIT7_TEACHER_FIXED)
    assert tp.find_general_high_source_skill_map("vocational", "數學1", _UNIT7) is None
    assert tp.find_general_high_source_skill_map("general", "數學2", _UNIT7) is None


def test_unit7_reimport_keeps_teacher_fixed_bindings(admin_app):
    """A re-import whose AI pass puts 進階題10/12 under Tangent again is corrected by the repository map only."""
    from models import TextbookExample, db

    titles = tp.find_general_high_source_skill_map("general", "數學1", _UNIT7)["titles"]
    ai_rows = [(t, _TANGENT if t in _UNIT7_TEACHER_FIXED else s) for t, s in titles.items()]
    ids = _seed_unit_rows(db, _UNIT7, ai_rows, [_JUDGE, _TANGENT])
    other = _seed_unit_rows(db, "單元6 圓方程式", [("7習題 進階題10", "gh_other_unit6")], ["gh_other_unit6"])

    summary = tp.apply_general_high_source_skill_map(
        parsed_data={"chapters": [{"chapter_title": _UNIT7}]},
        scan={"chapter": {}},
        curriculum_info={"curriculum": "general", "volume": "數學1"},
    )

    assert summary["applied"] is True and summary["map_id"] == "general_數學1_單元7_圓與直線的關係_v1"
    assert sorted(c["title"] for c in summary["changed"]) == sorted(_UNIT7_TEACHER_FIXED)
    assert all(c["from"] == _TANGENT and c["to"] == _JUDGE for c in summary["changed"])
    assert summary["unmapped"] == [] and summary["missing"] == []
    bound = {t: db.session.get(TextbookExample, i).skill_id for t, i in ids.items()}
    assert bound == titles
    assert sum(s == _JUDGE for s in bound.values()) == 21 and sum(s == _TANGENT for s in bound.values()) == 15
    assert db.session.get(TextbookExample, other["7習題 進階題10"]).skill_id == "gh_other_unit6"


def _fake_shared_enrichment(seen, attach_titles):
    import json

    def _fake(*, pdf_path, examples, curriculum_info, project_root=None, debug_dir=None, write_notes=True, publish_assets=None, **kw):
        seen.append({
            "project_root": project_root,
            "publish_assets": publish_assets,
            "curriculum_info": curriculum_info,
            "examples": list(examples),
            "assemble_visual_groups": kw.get("assemble_visual_groups"),
        })
        out = []
        for te in examples:
            notes = json.loads(te.notes or "{}")
            assert notes["question_anchor"]["anchor_id"]
            if te.source_description.split(" [")[0] in attach_titles:
                notes["image_assets"] = [{"path": "uploads/question_assets/x.png", "source": "pdf", "visual_status": "accepted"}]
                te.notes = json.dumps(notes, ensure_ascii=False)
                out.append({"id": te.id, "status": "mounted", "page": 2, "visual_bbox": [1, 2, 3, 4]})
            else:
                out.append({"id": te.id, "status": "skipped_low_confidence", "page": 2, "visual_bbox": [1, 2, 3, 4]})
        return {"ok": True, "mounted": len(attach_titles), "rows": out}

    return _fake


def test_general_high_enrichment_uses_temporary_asset_root_and_gates_review(admin_app, monkeypatch, tmp_path):
    import json

    import core.textbook_pdf_visual as visual_mod
    from models import TextbookExample, db

    app, _admin_id = admin_app
    missing = {"needs_review": True, "needs_review_before_visual": False, "visual_required_missing": True}
    rows = _seed_general_rows(
        db,
        [
            ("例題1", "已知 f(x)=x+1，求 f(2)。", {"needs_review": False}),
            ("例題2", "如右圖所示，求面積。", dict(missing)),
            ("隨堂練習2", "如圖所示，求長度。", dict(missing)),
        ],
    )
    ids = {title: te.id for title, te in rows.items()}
    seen = []
    monkeypatch.setattr(visual_mod, "enrich_textbook_examples_with_pdf_visuals", _fake_shared_enrichment(seen, {"例題2"}))
    pdf = tmp_path / "unit.pdf"
    pdf.write_bytes(b"%PDF-1.4\n")
    asset_root = tmp_path / "visual_assets"

    summary = tp.enrich_general_high_pdf_visuals(
        pdf_path=str(pdf),
        parsed_data={"chapters": [{"chapter_title": "單元3 函數"}]},
        scan=tp.scan_general_high_layout(GENERAL_HIGH_TEXT),
        curriculum_info={"curriculum": "general", "volume": "數學1", "publisher": "longteng"},
        import_policy={"visual_asset_root": str(asset_root)},
    )

    assert summary["ok"] is True
    assert seen[0]["project_root"] == str(asset_root)
    assert seen[0]["publish_assets"] is False
    assert seen[0]["assemble_visual_groups"] is True
    assert seen[0]["curriculum_info"]["chapter"] == "單元3 函數"
    assert len(seen[0]["examples"]) == 3
    notes = {title: json.loads(db.session.get(TextbookExample, i).notes) for title, i in ids.items()}
    assert notes["例題1"]["pdf_visual_status"] == tp.GH_NO_VISUAL_REQUIRED
    assert notes["例題1"]["needs_review"] is False
    assert notes["例題2"]["pdf_visual_status"] == tp.GH_VISUAL_ATTACHED
    assert notes["例題2"]["needs_review"] is False
    assert notes["例題2"]["visual_required_missing"] is False
    assert notes["隨堂練習2"]["pdf_visual_status"] == tp.GH_VISUAL_REQUIRED_MISSING
    assert notes["隨堂練習2"]["needs_review"] is True
    enrichment = notes["隨堂練習2"]["pdf_visual_enrichment"]
    assert enrichment["source_pdf"] == "unit.pdf" and enrichment["page"] == 2
    assert enrichment["visual_bbox"] == [1, 2, 3, 4]
    assert len(enrichment["source_pdf_sha256"]) == 64

    seen.clear()
    tp.enrich_general_high_pdf_visuals(
        pdf_path=str(pdf),
        parsed_data={"chapters": [{"chapter_title": "單元3 函數"}]},
        scan=tp.scan_general_high_layout(GENERAL_HIGH_TEXT),
        curriculum_info={"curriculum": "general", "volume": "數學1"},
    )
    assert seen[0]["project_root"] == app.root_path


def _synthetic_figure_page():
    # One figure split into three faces + a dimension bracket glyph + a short label,
    # a blended page-texture path touching it, and an unrelated figure far away.
    def d(bb, blended=False):
        return {"bbox": bb, "area": (bb[2] - bb[0]) * (bb[3] - bb[1]), "blended": blended}

    return {
        "page": 1,
        "width": 600.0,
        "height": 800.0,
        "drawings": [
            d([230.0, 216.0, 351.0, 273.0]),
            d([229.6, 241.0, 290.2, 300.0]),
            d([290.2, 241.0, 350.8, 300.1]),
            d([200.0, 280.0, 229.0, 330.0], blended=True),
            d([420.0, 520.0, 500.0, 600.0]),
        ],
        "small_drawings": [d([350.8, 241.0, 354.7, 252.2])],
        "images": [],
        "words": [
            (120.0, 180.0, 340.0, 194.0, "長方體紙盒，如下圖所示，求底面", 0, 0, 0),
            (356.0, 251.3, 382.0, 265.5, "2cm", 1, 0, 0),
        ],
    }


def test_visual_group_assembly_merges_adjacent_parts_and_short_label():
    from core.textbook_pdf_visual import assemble_visual_group_bbox

    page = _synthetic_figure_page()
    limit = [36.0, 170.0, 564.0, 460.0]
    grown = assemble_visual_group_bbox([290.2, 241.0, 350.8, 300.1], page, limit_bbox=limit)
    assert grown == [229.6, 216.0, 382.0, 300.1]
    assert assemble_visual_group_bbox([290.2, 241.0, 350.8, 300.1], page, limit_bbox=None) == [290.2, 241.0, 350.8, 300.1]


def test_classify_keeps_single_pick_unless_group_assembly_enabled():
    from core.textbook_pdf_visual import assign_question_regions, classify_and_detect_visuals

    def rows():
        return [{
            "id": 1,
            "problem_text": "長方體紙盒，如下圖所示，求底面。",
            "match_score": 0.99,
            "pdf_match": {"page": 1, "question_start_y": 180.0},
        }]

    page = _synthetic_figure_page()
    default = classify_and_detect_visuals(assign_question_regions(rows(), [page]), [page])[0]
    grouped = classify_and_detect_visuals(
        assign_question_regions(rows(), [page]), [page], assemble_visual_groups=True
    )[0]
    assert default["visual_bbox"] == [230.0, 216.0, 351.0, 273.0]
    assert default["should_mount"] is False
    assert grouped["visual_bbox"] == [229.6, 216.0, 382.0, 300.1]
    assert grouped["visual_classification"] == "required"
    assert grouped["should_mount"] is True


@pytest.mark.parametrize(
    "notes, text, expected, needs_review",
    [
        ({"needs_review": False}, "求 \\(1+1\\) 的值。", tp.GH_NO_VISUAL_REQUIRED, False),
        ({"needs_review": False}, "如右圖所示，求面積。", tp.GH_VISUAL_REQUIRED_MISSING, True),
        (
            {"needs_review": False, "image_assets": [{"path": "a.png", "visual_status": "needs_review"}]},
            "如圖，求長度。",
            tp.GH_VISUAL_UNCERTAIN,
            True,
        ),
        (
            {"needs_review": True, "needs_review_before_visual": False, "visual_required_missing": True,
             "image_warning": "missing_docx_image_asset", "image_assets": [{"path": "a.png", "visual_status": "accepted"}]},
            "如圖，求長度。",
            tp.GH_VISUAL_ATTACHED,
            False,
        ),
    ],
)
def test_required_visual_keeps_needs_review_until_attached(notes, text, expected, needs_review):
    status = tp.classify_general_high_visual_status(notes, text)
    assert status == expected
    out = tp._apply_general_high_visual_review(dict(notes), status)
    assert out["needs_review"] is needs_review
    assert out["pdf_visual_status"] == expected
    if expected == tp.GH_VISUAL_ATTACHED:
        assert "image_warning" not in out


# ---- /textbook_importer_v3 general V1 task status ----

GENERAL_IMPORT_MESSAGES = [
    "INFO: 正在從 u/01--word.docx 提取內容...",
    "INFO: [MATHTYPE CONVERT] ole=417 converted=417 failed=0 eq_fields=17 eq_converted=17",
    "INFO: [DOCX BLOCK SCAN] question_blocks=49",
    "INFO: --- 正在執行 AI 內容解析 ---",
    "INFO: 執行教材內容結構化解析作業 attempt 1/3",
    "INFO: [DOCX HYDRATE] filled_problem_text=49 scanned_blocks=49",
    "INFO: [IMPORT INVENTORY GUARD] returned_titles_count=49",
    "INFO: 正在將解析結果寫入資料庫...",
    "INFO: [PRACTICE IMPORT] detected title=隨堂練習1 source_type=in_class_practice linked_example=例題1 skill_id=gh_A",
    "INFO: [PRACTICE IMPORT] detected title=隨堂練習5 source_type=in_class_practice linked_example=例題3 skill_id=gh_B",
    "INFO: [PRACTICE IMPORT] detected title=隨堂練習7 source_type=in_class_practice linked_example=例題7 skill_id=gh_C",
    "INFO: [PRACTICE IMPORT] detected title=隨堂練習12 source_type=in_class_practice linked_example=例題12 skill_id=gh_D",
    "INFO: [PDF VISUAL] ok=True statuses={'NO_VISUAL_REQUIRED': 46, 'VISUAL_UNCERTAIN': 1, 'VISUAL_ATTACHED': 2} "
    "ambiguous=0 unanchored=0 warnings=[]",
    "INFO: Import complete: skills=0, curriculums=0, examples=15, practices=34, in_class=18, chapter_exercises=16, "
    "self_assessment=0, exam=0, other=0, needs_review=0, skipped=0",
]


def _run_general_v1_task(monkeypatch, fake_process, paths=("u/01--word.docx", "u/01--pdf.pdf")):
    import core.routes.admin as admin_mod
    from app import app
    from core.general_v1_import_status import get_general_v1_task_snapshot, register_general_v1_task

    task_id = f"t-{uuid.uuid4().hex[:8]}"
    q = register_general_v1_task(task_id)
    assert get_general_v1_task_snapshot(task_id)["state"] == "queued"
    monkeypatch.setattr(admin_mod.textbook_processor, "process_textbook_file", fake_process)
    admin_mod.background_processing(list(paths), q, app.app_context(), {"curriculum": "general"}, True)
    messages = []
    while not q.empty():
        messages.append(q.get())
    return task_id, get_general_v1_task_snapshot(task_id), messages


def test_general_v1_status_completed_with_real_counts(monkeypatch):
    from core.general_v1_import_status import STAGES

    seen_running = []

    def _fake_process(file_path, *, queue, **kwargs):
        for message in GENERAL_IMPORT_MESSAGES:
            queue.put(message)
            if "AI 內容解析" in message:
                seen_running.append(dict(queue.general_status.snapshot()))
        return {"status": "success"}

    _task_id, snap, messages = _run_general_v1_task(monkeypatch, _fake_process)
    running = seen_running[0]
    assert running["state"] == "running"
    assert running["current_stage"] == "question_parse"
    assert running["stages"]["ai_alignment"] == "running"
    assert running["stages"]["equation_convert"] == "done"

    assert snap["state"] == "completed" and snap["error"] is None
    assert snap["completed_stages"] == list(STAGES)
    r = snap["result"]
    assert (r["source_pairs"], r["successful_pairs"], r["failed_pairs"]) == (1, 1, 0)
    assert (r["parsed_questions"], r["imported_questions"]) == (49, 49)
    assert (r["formal_skills_created"], r["formal_skills_reused"]) == (0, 4)
    assert r["gemini_requests"] == 1
    assert r["visual_attached"] == 2 and r["visual_needs_review"] == 1 and r["needs_review"] == 1
    assert "INFO: [DOCX BLOCK SCAN] question_blocks=49" in messages
    assert messages[-1] == "END_OF_STREAM"


def test_general_v1_status_unreported_counts_stay_unknown(monkeypatch):
    def _fake_process(file_path, *, queue, **kwargs):
        queue.put("INFO: [DOCX BLOCK SCAN] question_blocks=3")
        return {"status": "success"}

    _task_id, snap, _messages = _run_general_v1_task(monkeypatch, _fake_process)
    r = snap["result"]
    assert snap["state"] == "completed"
    assert r["imported_questions"] is None
    assert r["formal_skills_reused"] is None
    assert r["gemini_requests"] is None
    assert snap["stages"]["equation_convert"] == "skipped"
    assert snap["stages"]["pdf_alignment"] == "pending"


@pytest.mark.parametrize("mode", ["error_result", "exception"])
def test_general_v1_status_failed_reports_stage_and_safe_error(monkeypatch, mode):
    def _fake_process(file_path, *, queue, **kwargs):
        queue.put(GENERAL_IMPORT_MESSAGES[0])
        queue.put(GENERAL_IMPORT_MESSAGES[1])
        queue.put(GENERAL_IMPORT_MESSAGES[2])
        if mode == "exception":
            raise RuntimeError("boom\nTraceback (most recent call last):\n  File secret.py")
        return {"status": "error", "message": "AI analysis failed."}

    _task_id, snap, _messages = _run_general_v1_task(monkeypatch, _fake_process)
    assert snap["state"] == "failed"
    assert snap["error"]["stage"] == "question_parse"
    assert snap["stages"]["question_parse"] == "failed"
    assert "Traceback" not in snap["error"]["message"]
    assert snap["error"]["message"] == ("boom" if mode == "exception" else "AI analysis failed.")
    assert snap["result"]["failed_pairs"] == 1


def test_general_v1_status_endpoint(admin_app, monkeypatch):
    def _fake_process(file_path, *, queue, **kwargs):
        for message in GENERAL_IMPORT_MESSAGES:
            queue.put(message)
        return {"status": "success"}

    task_id, _snap, _messages = _run_general_v1_task(monkeypatch, _fake_process)
    app, admin_id = admin_app
    client = app.test_client()
    with client.session_transaction() as sess:
        sess["_user_id"] = str(admin_id)
        sess["_fresh"] = True
    resp = client.get(f"/textbook_importer_v3/general_task/{task_id}")
    assert resp.status_code == 200
    data = resp.get_json()
    assert data["ok"] is True and data["pipeline"] == "general_v1" and data["state"] == "completed"
    assert client.get("/textbook_importer_v3/general_task/missing").status_code == 404


def test_vocational_v1_worker_queue_has_no_general_status(monkeypatch):
    import core.routes.admin as admin_mod
    from app import app

    q = queue_mod.Queue()
    monkeypatch.setattr(admin_mod.textbook_processor, "process_textbook_file", lambda *_a, **_k: {"status": "success"})
    admin_mod.background_processing(["u/1-1.docx", "u/1-1.pdf"], q, app.app_context(), {"curriculum": "vocational"}, True)
    assert not hasattr(q, "general_status")


def test_v3_template_general_polling_branch_and_stop_conditions():
    from pathlib import Path

    template = (Path(__file__).resolve().parents[1] / "templates" / "textbook_importer_v3.html").read_text(encoding="utf-8")
    branch = "if (result.data.pipeline === GENERAL_V1_PIPELINE && result.data.task_id && result.data.status_url) {"
    assert branch in template
    assert template.index(branch) < template.index("if (result.data.task_id && result.data.status_url) {")
    assert "if (data.state === 'completed') return 'completed';" in template
    assert "if (data.state === 'failed') return 'failed';" in template
    assert "if (decision === 'continue') {" in template
    assert "✅ 教材匯入完成" in template

