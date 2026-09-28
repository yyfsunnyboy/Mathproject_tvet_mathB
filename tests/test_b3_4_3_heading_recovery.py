# -*- coding: utf-8 -*-
"""B3 4-3 recovers one section-local numbered heading typo without editing the DOCX."""

from pathlib import Path

from docx import Document

from core.mathb_concept_heading import AUTHORITY_SOURCE_HEADING_RECOVERED
from core.textbook_importer_v3_docx import extract_docx_skill_headings
from core.textbook_importer_v3_pipeline import (
    STRUCTURAL_SKILL_NUMBERED_FOUND,
    STRUCTURAL_SKILL_PLAIN_FOUND,
    resolve_v3_curriculum_binding,
)

ROOT = Path(__file__).resolve().parents[1]
B3 = ROOT / "textbook_import" / "source" / "vocational" / "math_B3"
B2 = ROOT / "textbook_import" / "source" / "vocational" / "math_B2"


def _original_docx(folder: Path, fragment: str) -> Path:
    hits = [
        path
        for path in folder.glob("*.docx")
        if fragment in path.name and "_Latex" not in path.name
    ]
    assert len(hits) == 1, [path.name for path in hits]
    return hits[0]


def _bold(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True


def test_b3_4_3_recovers_section_local_heading_typo():
    path = _original_docx(B3, "4-3")
    before = path.read_bytes()
    audit = extract_docx_skill_headings(str(path), section_code="4-3", volume="數學B3")
    assert path.read_bytes() == before
    assert [item["concept_code"] for item in audit["skill_candidates"]] == ["4-3.1", "4-3.2"]
    assert audit["unresolved_heading_count"] == 0
    assert audit["plain_source_heading_count"] == 0
    first, second = audit["skill_candidates"]
    assert first["concept_name"] == "對數的意義"
    assert first["authority_source"] == "authoritative_numbered_concept_heading"
    assert first["printed_concept_code"] == "4-3.1"
    assert second["concept_name"] == "對數的性質"
    assert second["authority_source"] == AUTHORITY_SOURCE_HEADING_RECOVERED
    assert second["source_heading_raw"] == "3-3.2對數的性質"
    assert second["recovered_heading"] == "4-3.2對數的性質"
    assert second["printed_concept_code"] == "3-3.2"
    assert second["source_concept_code"] == "3-3.2"
    assert second["recovery_reason"] == (
        "section-local numbered heading typo; content and placement belong to 4-3"
    )
    decision = resolve_v3_curriculum_binding(
        section_heading=audit["section_heading"],
        same_section=True,
        candidate_count=audit["candidate_count"],
        unresolved_heading_count=audit["unresolved_heading_count"],
        outline_action="existing",
        chapter="第4章 指數與對數",
        section="4-3對數",
        section_code="4-3",
        plain_source_heading_count=0,
        plain_source_sequence_ok=None,
    )
    assert decision["structural_skill_status"] == STRUCTURAL_SKILL_NUMBERED_FOUND
    assert decision["curriculum_binding"] == "PASS"


def test_non_next_section_typo_stays_unresolved(tmp_path):
    path = tmp_path / "gap.docx"
    doc = Document()
    _bold(doc, "4-3對數")
    _bold(doc, "4-3.1對數的意義")
    _bold(doc, "3-3.4不是下一題")
    doc.save(path)
    audit = extract_docx_skill_headings(str(path), section_code="4-3", volume="數學B3")
    assert [item["concept_code"] for item in audit["skill_candidates"]] == ["4-3.1"]
    assert audit["unresolved_heading_count"] == 1


def test_existing_section_contracts_stay_in_place():
    numbered = extract_docx_skill_headings(
        str(_original_docx(B3, "3-1")), section_code="3-1", volume="數學B3"
    )
    assert [item["concept_code"] for item in numbered["skill_candidates"]] == ["3-1.1", "3-1.2", "3-1.3"]
    assert numbered["unresolved_heading_count"] == 0
    assert numbered["plain_source_heading_count"] == 0

    plain = extract_docx_skill_headings(
        str(_original_docx(B3, "3-2")), section_code="3-2", volume="數學B3"
    )
    assert plain["candidate_count"] == 0
    assert plain["unresolved_heading_count"] == 0
    assert plain["plain_source_heading_count"] > 0
    decision = resolve_v3_curriculum_binding(
        section_heading=plain["section_heading"],
        same_section=True,
        candidate_count=0,
        unresolved_heading_count=0,
        outline_action="existing",
        chapter="第3章 二元一次不等式及其應用",
        section="3-2二元一次不等式",
        section_code="3-2",
        plain_source_heading_count=plain["plain_source_heading_count"],
        plain_source_sequence_ok=plain["plain_source_sequence_ok"],
    )
    assert decision["structural_skill_status"] == STRUCTURAL_SKILL_PLAIN_FOUND

    chapter2 = extract_docx_skill_headings(
        str(_original_docx(B3, "2-1")), section_code="2-1", volume="數學B3"
    )
    assert chapter2["candidate_count"] > 0
    assert chapter2["unresolved_heading_count"] == 0
    assert chapter2["plain_source_heading_count"] == 0
    assert all(
        item["authority_source"] == "authoritative_numbered_concept_heading"
        for item in chapter2["skill_candidates"]
    )

    b2 = next(
        path for path in B2.glob("*4-2*_Latex.docx")
    )
    b2_audit = extract_docx_skill_headings(str(b2), section_code="4-2", volume="數學B2")
    assert [item["concept_code"] for item in b2_audit["skill_candidates"]] == [
        "4-2.1",
        "4-2.2",
        "4-2.3",
        "4-2.4",
    ]
    assert b2_audit["unresolved_heading_count"] == 0


def test_b3_4_3_questions_bind_by_heading_and_stem():
    from core.textbook_formal_concept import build_formal_skill_id_from_en_id
    import core.textbook_processor_v2 as processor
    from core.textbook_question_anchor import (
        build_anchors_from_block_meta,
        detect_anchor_id_collisions,
    )
    from core.textbook_structural_metadata import unique_logarithm_exercise_heading

    source = _original_docx(B3, "4-3")
    latex_hits = [
        path
        for path in B3.glob("*.docx")
        if "4-3" in path.name and path.name.endswith("_Latex.docx")
    ]
    assert len(latex_hits) == 1
    audit = extract_docx_skill_headings(str(source), section_code="4-3", volume="數學B3")
    info = {
        "curriculum": "vocational",
        "publisher": "longteng",
        "volume": "數學B3",
        "chapter": "第4章 指數與對數",
        "section": "4-3對數",
        "section_code": "4-3",
        "source_scope": "section_textbook",
        "structural_skill_candidates": audit["skill_candidates"],
    }
    lines = processor.phase1_extract_docx_lines(str(latex_hits[0]), curriculum_info=info)
    processor.phase2_deterministic_block_slice(
        lines, source_scope="section_textbook", curriculum_info=info, read_only=True
    )
    meta = dict(processor._DOCX_BLOCK_META)

    def skill_id(code: str) -> str:
        return build_formal_skill_id_from_en_id(
            volume="數學B3",
            concept_en_id=processor._fallback_en_id_from_concept_code(code),
        )

    bound = {}
    for title, block in meta.items():
        code = str(block.get("concept_code") or "")
        reason = "heading proximity"
        if not code:
            heading = unique_logarithm_exercise_heading(
                block.get("problem_text"), audit["skill_candidates"]
            )
            assert heading is not None, title
            code = heading["concept_code"]
            reason = heading["binding_reason"]
        bound[title] = (code, skill_id(code), reason)

    assert bound["例1"][0] == "4-3.1"
    assert bound["隨堂練習1"][0] == "4-3.1"
    for title in (
        "例2",
        "隨堂練習2",
        "例3",
        "隨堂練習3",
        "例4",
        "隨堂練習4",
        "例5",
        "隨堂練習5",
        "例6",
        "隨堂練習6",
        "例7",
        "隨堂練習7",
        "例8",
        "隨堂練習8",
        "113統測B",
    ):
        assert bound[title][0] == "4-3.2", title
        assert meta[title]["heading_provenance"]["source_heading_raw"] == "3-3.2對數的性質"
    assert bound["4-3習題 基礎題 1"][0] == "4-3.1"
    for number in range(2, 9):
        assert bound[f"4-3習題 基礎題 {number}"][0] == "4-3.2"
    assert bound["4-3習題 進階題 9"][0] == "4-3.2"
    assert bound["4-3習題 進階題 10"][0] == "4-3.2"
    assert len(bound) == 27
    assert all(skill.startswith("vh_數學B3_SubSection_4_3_") for _code, skill, _reason in bound.values())
    anchors = build_anchors_from_block_meta(meta, info)
    assert detect_anchor_id_collisions(anchors) == []
